"""Deterministic local file creation from validated plans and verified inputs.

This module belongs in the offline artifact runner. It has no provider, database,
HTTP client or model-code execution. Office preview rendering is a separate stage.
"""

import hashlib
import io
import math
import shutil
import sys
import zipfile
from pathlib import Path
from typing import Any

from hub.artifact_specs import (
    ArtifactSpec,
    DocumentPlan,
    FormulaCell,
    ImagePlan,
    PdfPlan,
    PresentationPlan,
    SpreadsheetPlan,
)
from hub.presentation_features import content as presentation_content
from hub.spreadsheet_features import cache_formula_results, decorate_sheet, native_formula
from hub.template_render import document as template_document
from hub.template_render import presentation as template_presentation

MAX_INPUT = 20 * 1024 * 1024
MAX_OUTPUT = 40 * 1024 * 1024


def verified_inputs(spec: ArtifactSpec, inputs: list[Path]) -> list[Path]:
    if len(inputs) != len(spec.sources) + bool(spec.template):
        raise ValueError("Quellenbindung unvollständig")
    source_inputs = inputs[: len(spec.sources)]
    for binding, path in zip(spec.sources, source_inputs, strict=True):
        if path.is_symlink() or any(parent.is_symlink() for parent in path.parents):
            raise ValueError("Verknüpfte Quelldatei nicht erlaubt")
        if not path.is_file() or not 0 < path.stat().st_size <= MAX_INPUT:
            raise ValueError("Ungültige Quellengröße")
        if hashlib.sha256(path.read_bytes()).hexdigest() != binding.sha256:
            raise ValueError("Quelldatei wurde verändert")
    if spec.template:
        path = inputs[-1]
        if path.is_symlink() or any(parent.is_symlink() for parent in path.parents):
            raise ValueError("Verknüpfte Vorlagendatei nicht erlaubt")
        if not path.is_file() or not 0 < path.stat().st_size <= MAX_INPUT:
            raise ValueError("Ungültige Vorlagengröße")
        if hashlib.sha256(path.read_bytes()).hexdigest() != spec.template.asset_sha256:
            raise ValueError("Vorlage wurde verändert")
    return source_inputs


def _document(plan: DocumentPlan, path: Path, inputs: list[Path]) -> None:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.85)
    normal = document.styles["Normal"]
    normal.font.name = "Liberation Sans"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0x18, 0x28, 0x42)
    normal.paragraph_format.space_after = Pt(8)
    document.add_heading(plan.title, 0)
    for item in plan.sections:
        document.add_heading(item.heading, 1)
        for text in item.paragraphs:
            document.add_paragraph(text)
        if item.table:
            table = document.add_table(rows=0, cols=len(item.table[0]))
            table.style = "Light Shading Accent 1"
            for row in item.table:
                for cell, text in zip(table.add_row().cells, row, strict=True):
                    cell.text = text
        for picture in item.images:
            clean = io.BytesIO()
            _read_image(inputs[picture.source]).save(clean, format="PNG")
            clean.seek(0)
            document.add_picture(clean, width=Inches(picture.width_inches))
            if picture.caption:
                document.add_paragraph(picture.caption, style="Caption")
    document.core_properties.title = plan.title
    document.core_properties.author = "obdaAI"
    document.save(str(path))
    check = Document(str(path))
    if check.paragraphs[0].text != plan.title:
        raise ValueError("Word-Dateiprüfung fehlgeschlagen")


def _presentation(plan: PresentationPlan, path: Path) -> None:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    for number, item in enumerate(plan.slides, 1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(0xF4, 0xF7, 0xFD)  # type: ignore[no-untyped-call]
        title = slide.shapes.add_textbox(Inches(0.7), Inches(0.5), Inches(11.9), Inches(1.2))
        title.text_frame.word_wrap = True
        paragraph = title.text_frame.paragraphs[0]
        paragraph.text = item.title
        paragraph.font.name = "Liberation Sans"
        paragraph.font.size = Pt(30)
        paragraph.font.bold = True
        paragraph.font.color.rgb = RGBColor(0x15, 0x39, 0x76)  # type: ignore[no-untyped-call]
        _fit_slide_text(title.text_frame, maximum=30, minimum=18, bold=True)
        if item.visual:
            for element in list(slide.shapes):
                element._element.getparent().remove(element._element)
            presentation_content(slide, item, number, plan.title, len(plan.slides))
            continue
        box = slide.shapes.add_textbox(Inches(0.9), Inches(2), Inches(11.5), Inches(4.6))
        box.text_frame.word_wrap = True
        for index, text in enumerate(item.bullets):
            paragraph = (
                box.text_frame.paragraphs[0] if index == 0 else box.text_frame.add_paragraph()
            )
            paragraph.text = "• " + text
            paragraph.font.name = "Liberation Sans"
            paragraph.font.size = Pt(22)
            paragraph.font.color.rgb = RGBColor(0x18, 0x28, 0x42)  # type: ignore[no-untyped-call]
            paragraph.space_after = Pt(18)
        _fit_slide_text(box.text_frame, maximum=22, minimum=16, spacing=18)
        footer = slide.shapes.add_textbox(Inches(0.8), Inches(7), Inches(11.8), Inches(0.3))
        paragraph = footer.text_frame.paragraphs[0]
        paragraph.text = f"{plan.title[:100]} · {number}/{len(plan.slides)}"
        paragraph.font.name = "Liberation Sans"
        paragraph.font.size = Pt(10)
        paragraph.font.color.rgb = RGBColor(0x52, 0x64, 0x83)  # type: ignore[no-untyped-call]
        _fit_slide_text(footer.text_frame, maximum=10, minimum=8)
        slide.notes_slide.notes_text_frame.text = item.notes
    deck.core_properties.title, deck.core_properties.author = plan.title, "obdaAI"
    deck.save(str(path))
    if len(Presentation(str(path)).slides) != len(plan.slides):
        raise ValueError("PowerPoint-Folienzahl stimmt nicht")


def _fit_slide_text(
    frame: Any, *, maximum: int, minimum: int, bold: bool = False, spacing: int = 0
) -> None:
    from pptx.util import Pt

    suffix = "Bold" if bold else "Regular"
    candidates = (
        [f"/System/Library/Fonts/Supplemental/Arial{' Bold' if bold else ''}.ttf"]
        if sys.platform == "darwin"
        else [
            f"/usr/share/fonts/truetype/liberation/LiberationSans-{suffix}.ttf",
            f"/usr/share/fonts/truetype/liberation2/LiberationSans-{suffix}.ttf",
        ]
    )
    font = next((path for path in candidates if Path(path).is_file()), None)
    if font is None:
        raise ValueError("Schrift für die Folienlayoutprüfung fehlt")
    # TextFitter measures glyph widths and line heights, but not paragraph spacing.
    # Reserve that space and retain a small safety margin for Office line wrapping.
    bottom = frame.margin_bottom
    frame.margin_bottom = bottom + Pt(spacing * len(frame.paragraphs) + 4)
    try:
        frame.fit_text(font_family="Liberation Sans", max_size=maximum, bold=bold, font_file=font)
    except (TypeError, ValueError):
        raise ValueError("Folientext passt nicht in das Layout; Text kürzen") from None
    finally:
        frame.margin_bottom = bottom
    sizes = [run.font.size.pt for paragraph in frame.paragraphs for run in paragraph.runs]
    if not sizes or min(sizes) < minimum:
        raise ValueError("Folientext wäre zu klein; Text kürzen")
    for paragraph in frame.paragraphs:
        paragraph.font.size = Pt(min(sizes))


def _pdf(plan: PdfPlan, inputs: list[Path], path: Path) -> None:
    from pypdf import PdfReader, PdfWriter

    if plan.sections:
        from hub.artifact_preview import office_pdf

        temporary = path.with_suffix(".docx")
        try:
            _document(
                DocumentPlan(kind="docx", title=plan.title, sections=plan.sections),
                temporary,
                inputs,
            )
            office_pdf(temporary, path.parent)
        finally:
            temporary.unlink(missing_ok=True)
        return

    readers = [PdfReader(str(source), strict=True) for source in inputs]
    for reader in readers:
        if reader.is_encrypted:
            raise ValueError("Geschützte PDF vor der Bearbeitung entsperren")
        # A modified signed document would no longer validate. Require a separate
        # explicitly reviewed policy rather than silently invalidating signatures.
        fields = reader.get_fields() or {}
        if any(field.get("/FT") == "/Sig" for field in fields.values()):
            raise ValueError("Signierte PDF kann hier nicht verändert werden")
    writer = PdfWriter()
    for selection in plan.pages:
        reader = readers[selection.source]
        if selection.page > len(reader.pages):
            raise ValueError("PDF-Seite existiert nicht")
        page = writer.add_page(
            reader.pages[selection.page - 1], excluded_keys=["/AA", "/Annots", "/B"]
        )
        if selection.rotation:
            page.rotate(selection.rotation)
    writer.add_metadata({"/Title": plan.title, "/Producer": "obdaAI"})
    with path.open("wb") as stream:
        writer.write(stream)
    if len(PdfReader(str(path), strict=True).pages) != len(plan.pages):
        raise ValueError("PDF-Seitenzahl stimmt nicht")


def _read_image(path: Path) -> Any:
    from PIL import Image, ImageOps

    with Image.open(path) as original:
        if original.width * original.height > 24_000_000:
            raise ValueError("Quelldatei überschreitet das Pixellimit")
        if original.format not in {"PNG", "JPEG", "WEBP"}:
            raise ValueError("Bildformat nicht freigegeben")
        # Apply orientation before removing metadata. A new pixel-only image also
        # drops EXIF/GPS, comments and profiles when embedded in Word or text PDFs.
        transposed = ImageOps.exif_transpose(original)
        alpha = "A" in transposed.getbands() or "transparency" in transposed.info
        oriented = transposed.convert("RGBA" if alpha else "RGB")
        clean = Image.new(oriented.mode, oriented.size)
        clean.paste(oriented)
        return clean


def _image(plan: ImagePlan, inputs: list[Path], path: Path) -> None:
    from PIL import Image, ImageDraw, ImageFont, ImageOps

    rows = math.ceil(len(plan.tiles) / plan.columns)
    canvas = Image.new("RGB", (plan.columns * plan.tile_width, rows * plan.tile_height), "#f4f7fd")
    for number, tile in enumerate(plan.tiles):
        image = _read_image(inputs[tile.source])
        left, top, right, bottom = tile.crop
        crop = (
            int(left * image.width),
            int(top * image.height),
            int(right * image.width),
            int(bottom * image.height),
        )
        if crop[2] <= crop[0] or crop[3] <= crop[1]:
            raise ValueError("Zuschnitt enthält keine Pixel")
        font: Any = None
        if tile.caption:
            fonts = (
                ["/System/Library/Fonts/Supplemental/Arial.ttf"]
                if sys.platform == "darwin"
                else [
                    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
                    "/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf",
                    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
                ]
            )
            for candidate in fonts:
                if Path(candidate).is_file():
                    font = ImageFont.truetype(candidate, size=16)
                    break
            if font is None:
                raise ValueError("Beschriftung benötigt die installierte lokale Unicode-Schrift")
            missing = bytes(font.getmask("\U0010ffff"))
            if any(bytes(font.getmask(char)) == missing for char in set(tile.caption)):
                raise ValueError("Die lokale Schrift unterstützt nicht alle Beschriftungszeichen")
        lines: list[str] = []
        draw = ImageDraw.Draw(canvas)
        for word in tile.caption.split():
            if (
                not lines
                or draw.textlength(lines[-1] + " " + word, font=font) > plan.tile_width - 16
            ):
                lines.append(word)
            else:
                lines[-1] += " " + word
        if any(draw.textlength(line, font=font) > plan.tile_width - 16 for line in lines):
            raise ValueError("Beschriftung enthält ein zu langes Wort")
        band = 12 + len(lines) * 21 if lines else 0
        if band > plan.tile_height // 2:
            raise ValueError("Beschriftung benötigt mehr Platz; Bildhöhe erhöhen")
        fitted = ImageOps.fit(image.crop(crop), (plan.tile_width, plan.tile_height - band))
        x = (number % plan.columns) * plan.tile_width
        y = (number // plan.columns) * plan.tile_height
        canvas.paste(
            fitted,
            (x, y),
            fitted.getchannel("A") if fitted.mode == "RGBA" else None,
        )
        for line_number, line in enumerate(lines):
            draw.text(
                (x + 8, y + plan.tile_height - band + 6 + line_number * 21),
                line,
                font=font,
                fill="#182842",
            )
    canvas.save(path, format="PNG" if plan.format == "png" else "JPEG")
    with Image.open(path) as check:
        check.verify()


def _spreadsheet(plan: SpreadsheetPlan, path: Path) -> None:
    from openpyxl import Workbook, load_workbook  # type: ignore[import-untyped]
    from openpyxl.styles import (  # type: ignore[import-untyped]
        Alignment,
        Border,
        Font,
        PatternFill,
        Side,
    )
    from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]

    book = Workbook()
    book.remove(book.active)
    for item in plan.sheets:
        sheet = book.create_sheet(item.name)
        sheet.append(item.columns)
        for row in item.rows:
            sheet.append([
                native_formula(cell.formula) if isinstance(cell, FormulaCell) else cell
                for cell in row
            ])
        if item.totals:
            functions = {"sum": 109, "average": 101, "count": 103, "min": 105, "max": 104}
            sheet.append(
                [
                    f"=SUBTOTAL({functions[item.totals[column]]},"
                    f"{column}2:{column}{len(item.rows) + 1})"
                    if column in item.totals
                    else ""
                    for column in (
                        get_column_letter(number) for number in range(1, len(item.columns) + 1)
                    )
                ]
            )
        for row in sheet:
            for cell in row:
                # Model strings remain literal cells, never formulas/external links.
                original = (
                    item.rows[cell.row - 2][cell.column - 1]
                    if 1 < cell.row <= len(item.rows) + 1
                    else None
                )
                total_formula = cell.row == len(item.rows) + 2 and cell.column_letter in item.totals
                if (
                    isinstance(cell.value, str)
                    and not hasattr(original, "formula")
                    and not total_formula
                ):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.font = Font(name="Liberation Sans", size=11, color="020623")
        for cell in sheet[1]:
            cell.border = Border(bottom=Side(style="medium", color="E6531F"))
            cell.fill = PatternFill("solid", fgColor="020623")
            cell.font = Font(name="Liberation Sans", size=11, bold=True, color="FFFFFF")
        for index, column in enumerate(item.columns, 1):
            sheet.column_dimensions[get_column_letter(index)].width = min(
                48, max(16, len(column) + 4)
            )
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = "1:1"
        decorate_sheet(sheet, item, len(book.worksheets))
    book.properties.title, book.properties.creator = plan.title, "obdaAI"
    book.save(str(path))
    cache_formula_results(path)
    check = load_workbook(path, read_only=True, data_only=False)
    try:
        if check.sheetnames != [item.name for item in plan.sheets]:
            raise ValueError("Excel-Dateiprüfung fehlgeschlagen")
    finally:
        check.close()


def render(spec: ArtifactSpec, inputs: list[Path], output: Path) -> dict[str, Any]:
    """Write into a fresh private directory, never overwrite an existing version.

    The caller supplies only server-bound input paths and a new job output directory.
    No title, model parameter or source value becomes a filesystem path.
    """
    source_inputs = verified_inputs(spec, inputs)
    if output.exists() or output.is_symlink() or any(p.is_symlink() for p in output.parents):
        raise ValueError("Ausgabeversion existiert bereits")
    output.mkdir(mode=0o700)
    plan = spec.plan
    split = isinstance(plan, PdfPlan) and bool(plan.split_after)
    suffix = "zip" if split else plan.format if isinstance(plan, ImagePlan) else plan.kind
    pieces = []
    path = output / ("artifact." + suffix)
    try:
        if isinstance(plan, DocumentPlan):
            if spec.template:
                template_document(plan, path, inputs[-1], source_inputs)
            else:
                _document(plan, path, source_inputs)
        elif isinstance(plan, PresentationPlan):
            if spec.template:
                template_presentation(plan, path, inputs[-1], spec.template.style == "modern")
            else:
                _presentation(plan, path)
        elif isinstance(plan, SpreadsheetPlan):
            _spreadsheet(plan, path)
        elif isinstance(plan, PdfPlan):
            if split:
                # The joined PDF is only the ordered review surface. Downloads
                # consist of these exact, separately checked immutable PDFs.
                _pdf(plan, inputs, output / "artifact.pdf")
                folder = output / "parts"
                folder.mkdir(mode=0o700)
                edges = [0, *plan.split_after, len(plan.pages)]
                with zipfile.ZipFile(path, "x", compression=zipfile.ZIP_STORED) as archive:
                    total = 0
                    for number, (start, end) in enumerate(
                        zip(edges[:-1], edges[1:], strict=True), 1
                    ):
                        part = folder / ("part-" + str(number) + ".pdf")
                        _pdf(
                            PdfPlan(kind="pdf", title=plan.title, pages=plan.pages[start:end]),
                            inputs,
                            part,
                        )
                        total += part.stat().st_size
                        if total > MAX_OUTPUT:
                            raise ValueError("PDF-Teile überschreiten das Ausgabegrößenlimit")
                        content = part.read_bytes()
                        archive.writestr(part.name, content)
                        part.chmod(0o400)
                        pieces.append(
                            {
                                "filename": "parts/" + part.name,
                                "sha256": hashlib.sha256(content).hexdigest(),
                                "bytes": len(content),
                                "first_page": start + 1,
                                "pages": end - start,
                            }
                        )
                (output / "artifact.pdf").chmod(0o400)
            else:
                _pdf(plan, inputs, path)
        else:
            _image(plan, inputs, path)
        if not 0 < path.stat().st_size <= MAX_OUTPUT:
            raise ValueError("Ausgabe überschreitet das Größenlimit")
        path.chmod(0o400)
        return {
            "filename": path.name,
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "bytes": path.stat().st_size,
            "label": "LOCAL_ONLY",
            "preview_verified": False,
            **({"parts": pieces} if split else {}),
        }
    except Exception:
        shutil.rmtree(output)
        raise
