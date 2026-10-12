"""Bounded deterministic parser subprocess. Never interpret document instructions."""

import csv
import io
import json
import os
import resource
import subprocess
import sys
import tempfile
import warnings
import zipfile
from pathlib import Path
from typing import Any

from hub.extraction_errors import ExtractionError, classify_error

# Load fixed application code before Landlock; editable installs may live outside sys.prefix.
from hub.ocr_local import needs_ocr, prepare, raster_edge

# Original-file size is independent of decompressed content, text and page limits.
MAX_BYTES = 40 * 1024 * 1024
MAX_TEXT = 800_000
MAX_UNITS = 2000
EXTENSIONS = {
    ".txt",
    ".md",
    ".csv",
    ".docx",
    ".pptx",
    ".xlsx",
    ".pdf",
    ".png",
    ".jpg",
    ".jpeg",
    ".webp",
    ".tif",
    ".tiff",
}


def safe_archive(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        members = archive.infolist()
        if len(members) > 3000 or sum(m.file_size for m in members) > 80 * 1024 * 1024:
            raise ValueError("Archivgrenze überschritten")
        for member in members:
            name = member.filename.lower()
            if (
                member.flag_bits & 1
                or member.file_size > 20 * 1024 * 1024
                or member.file_size > max(1, member.compress_size) * 200
                or ".." in Path(name).parts
                or name.startswith("/")
                or "vbaproject" in name
                or "/embeddings/" in name
                or name.endswith((".exe", ".js", ".vbs"))
            ):
                raise ValueError("Aktive, verschlüsselte oder gefährliche Archiv-Inhalte")


def ocr(path: Path, cwd: Path, timeout: int = 30) -> str:
    prepared = cwd / "ocr-prepared.png"
    prepare(path, prepared)
    # One CPU thread prevents OpenMP oversubscription and memory spikes within
    # the parser's existing sandbox. PSM 1 adds right-angle orientation detection.
    env = {**os.environ, "OMP_THREAD_LIMIT": "1", "OMP_NUM_THREADS": "1"}
    try:
        result = subprocess.run(
            ["tesseract", str(prepared), "stdout", "-l", "deu+eng", "--dpi", "300", "--psm", "1"],
            cwd=cwd,
            env=env,
            capture_output=True,
            timeout=timeout,
            check=True,
        )
    except FileNotFoundError:
        raise ExtractionError("OCR_RUNTIME_MISSING") from None
    except subprocess.TimeoutExpired:
        raise ExtractionError("OCR_TIMEOUT") from None
    except subprocess.CalledProcessError as error:
        code = (
            "OCR_RUNTIME_MISSING"
            if b"Failed loading language" in (error.stderr or b"")
            else "OCR_FAILED"
        )
        raise ExtractionError(code) from None
    if len(result.stdout) > MAX_TEXT:
        raise ValueError("OCR-Ausgabe zu groß")
    return result.stdout.decode("utf-8")


def parse(path: Path, suffix: str) -> dict[str, Any]:
    if suffix not in EXTENSIONS or path.stat().st_size > MAX_BYTES:
        raise ValueError("Dateityp oder Dateigröße nicht unterstützt")
    chunks: list[dict[str, str]] = []
    notes: list[str] = []

    def add(locator: str, text: str, structured: str = "") -> None:
        if (
            len(chunks) >= MAX_UNITS
            or sum(len(c["content"]) for c in chunks) + len(text) > MAX_TEXT
        ):
            raise ValueError("Extraktionsgrenze überschritten")
        if text.strip():
            chunks.append({"locator": locator[:240], "content": text, "structured": structured})

    if suffix in {".docx", ".pptx", ".xlsx"}:
        safe_archive(path)
    if suffix in {".md", ".txt"}:
        content = path.read_text(encoding="utf-8-sig")
        if len(content) > MAX_TEXT:
            raise ValueError("Extraktionsgrenze überschritten")
        lines = content.splitlines()
        heading = "Text"
        if len(lines) <= MAX_UNITS:
            # Preserve existing small-file locators and immutable receipts.
            for i, line in enumerate(lines, 1):
                if line.startswith("#"):
                    heading = line.lstrip("# ")[:100]
                add(f"{heading} · Zeile {i}", line)
        else:
            # Line count alone is not document complexity. Bounded contiguous
            # windows retain every line and heading with exact source ranges.
            start = 1
            window: list[str] = []
            used = 0
            for i, line in enumerate(lines, 1):
                if window and (line.startswith("#") or used + len(line) + 1 > 1600):
                    add(f"{heading} · Zeilen {start}–{i - 1}", "\n".join(window))
                    window, used = [], 0
                if not window:
                    start = i
                if line.startswith("#"):
                    heading = line.lstrip("# ")[:100]
                window.append(line)
                used += len(line) + 1
            if window:
                add(f"{heading} · Zeilen {start}–{len(lines)}", "\n".join(window))
    elif suffix == ".csv":
        content = path.read_text(encoding="utf-8-sig")
        try:
            dialect = csv.Sniffer().sniff(content[:8192], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        for i, row in enumerate(csv.reader(io.StringIO(content), dialect), 1):
            add(f"CSV Zeile {i}", " | ".join(row), json.dumps(row, ensure_ascii=False))
    elif suffix == ".docx":
        from docx import Document

        doc = Document(str(path))
        for i, paragraph in enumerate(doc.paragraphs, 1):
            add(f"Dokumentabschnitt {i}", paragraph.text)
        for i, table in enumerate(doc.tables, 1):
            for j, row in enumerate(table.rows, 1):
                cells = [c.text for c in row.cells]
                add(f"Tabelle {i} · Zeile {j}", " | ".join(cells), json.dumps(cells))
        notes.append("Abschnitte statt Seiten; eingebettete Bilder/Diagramme nicht interpretiert")
    elif suffix == ".pptx":
        from pptx import Presentation

        for i, slide in enumerate(Presentation(str(path)).slides, 1):
            for shape in slide.shapes:
                if shape.has_text_frame:
                    add(f"Folie {i}", shape.text)
                if shape.has_table:
                    for j, row in enumerate(shape.table.rows, 1):
                        cells = [c.text for c in row.cells]
                        add(f"Folie {i} · Tabelle Zeile {j}", " | ".join(cells), json.dumps(cells))
        notes.append("Folientext und Tabellen; Bilder/Diagramme nicht interpretiert")
    elif suffix == ".xlsx":
        from openpyxl import load_workbook  # type: ignore[import-untyped]
        from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]

        wb = load_workbook(
            io.BytesIO(path.read_bytes()), read_only=True, data_only=False, keep_links=False
        )
        try:
            for sheet in wb:
                if sheet.max_row > MAX_UNITS or sheet.max_column > 200:
                    raise ValueError("Tabellenlimits überschritten")
                for i, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    cells = ["" if c is None else str(c) for c in row]
                    add(
                        f"{sheet.title}!A{i}:{get_column_letter(len(cells))}{i}",
                        " | ".join(cells),
                        json.dumps(cells),
                    )
        finally:
            wb.close()
        notes.append("Formeln erhalten, nicht berechnet; keine externe Verknüpfung ausgeführt")
    elif suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path, strict=True)
        if reader.is_encrypted:
            raise ValueError("Passwortgeschütztes PDF")
        if len(reader.pages) > 100:
            raise ValueError("PDF-Seitenlimit überschritten")
        for i, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            if needs_ocr(text, has_images=bool(page.images)):
                with tempfile.TemporaryDirectory() as tmp:
                    target = Path(tmp)
                    try:
                        subprocess.run(
                            [
                                "pdftoppm",
                                "-f",
                                str(i),
                                "-l",
                                str(i),
                                "-r",
                                "300",
                                "-scale-to",
                                str(
                                    raster_edge(
                                        float(page.mediabox.width), float(page.mediabox.height)
                                    )
                                ),
                                "-singlefile",
                                "-png",
                                str(path),
                                str(target / "page"),
                            ],
                            check=True,
                            capture_output=True,
                            timeout=30,
                        )
                    except FileNotFoundError:
                        raise ExtractionError("OCR_RUNTIME_MISSING") from None
                    except subprocess.TimeoutExpired:
                        raise ExtractionError("OCR_TIMEOUT") from None
                    except subprocess.CalledProcessError:
                        raise ExtractionError("OCR_FAILED") from None
                    recognized = ocr(target / "page.png", target, 120)
                    if recognized.strip():
                        text = recognized
                notes.append(
                    f"Seite {i}: lokale OCR (bis 300 dpi, Ausrichtung und Kontrast); "
                    "Erkennung kann fehlerhaft sein"
                )
            if not text.strip():
                text = (
                    f"PDF-Seite {i}: {float(page.mediabox.width):g} × "
                    f"{float(page.mediabox.height):g} Punkte. Kein Text extrahiert "
                    "oder durch lokale OCR erkannt. Seiteninhalt wurde nicht semantisch bewertet."
                )
            add(f"PDF Seite {i}", text)
    else:
        from PIL import Image

        Image.MAX_IMAGE_PIXELS = 20_000_000
        with warnings.catch_warnings():
            warnings.simplefilter("error")
            with Image.open(path) as image:
                image.verify()
            with Image.open(path) as image:
                with tempfile.TemporaryDirectory() as tmp:
                    target = Path(tmp)
                    text = ocr(path, target)
                    if text.strip():
                        add("Bild · lokale OCR", text)
                    else:
                        add(
                            "Bild · technische Daten",
                            f"Bilddatei {image.format}: {image.width} × {image.height} Pixel. "
                            "Kein Text durch lokale OCR erkannt. Bildinhalt und Diagrammsemantik "
                            "wurden nicht erkannt oder bewertet.",
                        )
        notes.append("Lokale OCR liest Text; keine zuverlässige Bild-/Diagrammsemantik")
    if not chunks:
        raise ValueError("Keine lesbaren Inhalte extrahiert")
    return {"chunks": chunks, "notes": list(dict.fromkeys(notes))}


def isolated_parse(path: Path, suffix: str) -> dict[str, Any]:
    # Child receives no DB URL, keys, provider configuration or inherited proxies.
    try:
        with tempfile.TemporaryDirectory(prefix="hub-parser-") as work:
            result = subprocess.run(
                [sys.executable, "-m", "hub.extract", str(path.resolve()), suffix],
                env={
                    "PATH": "/usr/bin:/bin",
                    "PYTHONDONTWRITEBYTECODE": "1",
                    "HUB_PARSER_WORKDIR": work,
                    "TMPDIR": work,
                    "HOME": work,
                },
                cwd=work,
                capture_output=True,
                timeout=600 if suffix == ".pdf" else 120,
            )
    except subprocess.TimeoutExpired:
        raise ExtractionError("PARSER_TIMEOUT") from None
    if result.returncode:
        raise ExtractionError(
            "PARSER_RESOURCE_LIMIT" if result.returncode < 0 else "EXTRACTION_FAILED"
        )
    if len(result.stdout) > 4_000_000:
        raise ExtractionError("TEXT_LIMIT")
    try:
        decoded = json.loads(result.stdout)
        if not isinstance(decoded, dict):
            raise ValueError
        if "error" in decoded:
            raise ExtractionError(
                decoded["error"] if isinstance(decoded["error"], str) else "EXTRACTION_FAILED"
            )
        chunks, notes = decoded["chunks"], decoded["notes"]
        if not isinstance(chunks, list) or not isinstance(notes, list) or not chunks:
            raise ValueError
        if len(chunks) > MAX_UNITS or any(not isinstance(note, str) for note in notes):
            raise ValueError
        total = 0
        for chunk in chunks:
            if not isinstance(chunk, dict) or set(chunk) != {"locator", "content", "structured"}:
                raise ValueError
            if (
                any(not isinstance(value, str) for value in chunk.values())
                or len(chunk["locator"]) > 240
            ):
                raise ValueError
            total += len(chunk["content"])
        if total > MAX_TEXT:
            raise ExtractionError("TEXT_LIMIT")
        return decoded
    except (ValueError, KeyError, TypeError) as error:
        raise classify_error(error) from None


if __name__ == "__main__":
    suffix = sys.argv[2]
    resource.setrlimit(resource.RLIMIT_CPU, (480, 480) if suffix == ".pdf" else (90, 90))
    resource.setrlimit(resource.RLIMIT_FSIZE, (32 * 1024 * 1024, 32 * 1024 * 1024))
    resource.setrlimit(resource.RLIMIT_NOFILE, (64, 64))
    if sys.platform == "linux":
        resource.setrlimit(resource.RLIMIT_AS, (768 * 1024 * 1024, 768 * 1024 * 1024))
    from hub.parser_sandbox import restrict

    source = Path(sys.argv[1])
    workdir = Path(os.environ["HUB_PARSER_WORKDIR"])
    try:
        restrict(source, workdir)
        output = parse(source, suffix)
    except Exception as error:
        output = {"error": classify_error(error, suffix).code}
    print(json.dumps(output, ensure_ascii=False))
