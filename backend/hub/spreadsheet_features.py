"""Finite spreadsheet data capabilities; no code, external references or macros."""

from __future__ import annotations

import math
import re
import subprocess
import sys
import tempfile
import zipfile
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import TYPE_CHECKING, Any
from xml.etree import ElementTree as ET

from openpyxl import load_workbook  # type: ignore[import-untyped]
from openpyxl.chart import (  # type: ignore[import-untyped]
    AreaChart,
    BarChart,
    LineChart,
    PieChart,
    Reference,
    ScatterChart,
    Series,
)
from openpyxl.chart.data_source import StrRef  # type: ignore[import-untyped]
from openpyxl.chart.series import SeriesLabel  # type: ignore[import-untyped]
from openpyxl.formula.tokenizer import Tokenizer  # type: ignore[import-untyped]
from openpyxl.styles import Font, PatternFill  # type: ignore[import-untyped]
from openpyxl.utils.cell import range_boundaries  # type: ignore[import-untyped]
from openpyxl.utils.datetime import (  # type: ignore[import-untyped]
    time_to_days,
    timedelta_to_days,
    to_excel,
)
from openpyxl.worksheet.filters import AutoFilter  # type: ignore[import-untyped]
from openpyxl.worksheet.table import (  # type: ignore[import-untyped]
    Table,
    TableColumn,
    TableStyleInfo,
)

if TYPE_CHECKING:
    from .artifact_specs import SpreadsheetPlan

FUNCTIONS = frozenset(
    (
        "SUM AVERAGE MIN MAX COUNT COUNTA COUNTBLANK IF IFS IFERROR AND OR NOT "
        "ROUND ROUNDUP ROUNDDOWN ABS INT MOD SUMIF SUMIFS COUNTIF COUNTIFS "
        "AVERAGEIF AVERAGEIFS SUMPRODUCT VLOOKUP HLOOKUP INDEX MATCH XLOOKUP "
        "SUBTOTAL CONCAT CONCATENATE LEFT RIGHT MID LEN TRIM UPPER LOWER TEXT VALUE "
        "DATE YEAR MONTH DAY EOMONTH TODAY NPV IRR PMT POWER SQRT MEDIAN LARGE SMALL"
    ).split()
)
FUTURE_FUNCTIONS = frozenset({"XLOOKUP", "IFS", "CONCAT"})


def native_formula(formula: str) -> str:
    """Emit the MS-XLSX future-function prefix without changing plan receipts.

    Token-level changes never rewrite text literals or sheet names that happen
    to contain a function name. Only previously validated functions are mapped.
    """
    return "=" + "".join(
        "_xlfn." + token.value.upper()
        if token.type == "FUNC"
        and token.subtype == "OPEN"
        and token.value[:-1].upper() in FUTURE_FUNCTIONS
        else token.value
        for token in Tokenizer(formula).items
    )


def _function_arguments(tokens: list[Any], start: int) -> tuple[list[list[Any]], int]:
    """Split one tokenized function without splitting nested expressions or text."""
    arguments: list[list[Any]] = [[]]
    depth = 1
    for index in range(start + 1, len(tokens)):
        token = tokens[index]
        if token.type in {"FUNC", "PAREN"}:
            if token.subtype == "OPEN":
                depth += 1
            else:
                depth -= 1
                if depth == 0:
                    return arguments, index
        if token.type == "SEP" and token.subtype == "ARG" and depth == 1:
            arguments.append([])
        else:
            arguments[-1].append(token)
    raise ValueError("Formelklammern sind unvollständig")


def calculation_formula(formula: str) -> str:
    """Lower exact-match XLOOKUP for offline Calc 7.4, only in the calculation copy.

    The downloadable workbook keeps native _xlfn.XLOOKUP and its calculated cache.
    Validated one-dimensional arrays permit equivalent INDEX/MATCH evaluation.
    """

    def expression(tokens: list[Any]) -> str:
        result = ""
        index = 0
        while index < len(tokens):
            token = tokens[index]
            if token.type == "FUNC" and token.subtype == "OPEN":
                arguments, end = _function_arguments(tokens, index)
                rendered = [expression(argument) for argument in arguments]
                if token.value[:-1].split(".")[-1].upper() == "XLOOKUP":
                    if len(rendered) not in {3, 4}:
                        raise ValueError(
                            "XLOOKUP unterstützt genaues Suchen mit drei oder vier Argumenten"
                        )
                    lookup, lookup_range, return_range = rendered[:3]
                    escaped = (
                        f'SUBSTITUTE(SUBSTITUTE(SUBSTITUTE({lookup},"~","~~"),"*","~*"),"?","~?")'
                    )
                    exact_lookup = f"IF(ISTEXT({lookup}),{escaped},{lookup})"
                    match = f"MATCH({exact_lookup},{lookup_range},0)"
                    reference = _REF.fullmatch(return_range.strip())
                    if reference is None:
                        raise ValueError("XLOOKUP benötigt einen begrenzten Ergebnisbereich")
                    left, top, right, bottom = range_boundaries(
                        reference[3] + ":" + (reference[4] or reference[3])
                    )
                    value = (
                        f"INDEX({return_range},1,{match})"
                        if top == bottom and left != right
                        else f"INDEX({return_range},{match},1)"
                    )
                    result += (
                        f"IF(ISNA({match}),{rendered[3]},{value})" if len(rendered) == 4 else value
                    )
                else:
                    result += token.value + ",".join(rendered) + ")"
                index = end + 1
            else:
                result += token.value
                index += 1
            if len(result) > 12_000:
                raise ValueError("Excel-Berechnungsausdruck überschreitet das Limit")
        return result

    return "=" + expression(Tokenizer(formula).items)


FORMATS = {
    "currency": '_("€"* #,##0_);_("€"* \\(#,##0\\);_("€"* "-"_);_(@_)',
    "percent": '0.0%;(0.0%);"-"',
    "integer": '#,##0;(#,##0);"-"',
    "decimal": '#,##0.00;(#,##0.00);"-"',
    "date": "dd-mmm-yy",
}
_REF = re.compile(
    r"(?:(?:'((?:[^']|'')+)'|([A-Za-z_][A-Za-z0-9_ ]*))!)?"
    r"(\$?[A-T]\$?[1-9][0-9]{0,2})(?::(\$?[A-T]\$?[1-9][0-9]{0,2}))?$"
)


def bounded_reference(
    value: str, current: str, dimensions: dict[str, tuple[int, int]]
) -> tuple[str, int, int, int, int]:
    match = _REF.fullmatch(value)
    if not match:
        raise ValueError("Nur begrenzte A1-Zellverweise ohne externe Verbindungen erlaubt")
    name = (match[1].replace("''", "'") if match[1] else match[2]) or current
    canonical = next((sheet for sheet in dimensions if sheet.casefold() == name.casefold()), None)
    if canonical is None:
        raise ValueError("Formel verweist auf unbekanntes Tabellenblatt")
    start, end = match[3], match[4] or match[3]
    left, top, right, bottom = range_boundaries(start + ":" + end)
    columns, rows = dimensions[canonical]
    if not (1 <= left <= right <= columns and 1 <= top <= bottom <= rows):
        raise ValueError("Zellverweis außerhalb der vorgesehenen Tabelle")
    return canonical, left, top, right, bottom


def formula_references(
    formula: str, current: str, dimensions: dict[str, tuple[int, int]], *, strict: bool = True
) -> set[tuple[str, int, int]]:
    if not formula.startswith("=") or len(formula) > 2000:
        raise ValueError("Formel muss mit = beginnen und begrenzt sein")
    try:
        tokens = Tokenizer(formula).items
    except Exception as exc:
        raise ValueError("Ungültige Excel-Formel") from exc
    if not tokens or len(tokens) > 500:
        raise ValueError("Formel ist leer oder zu komplex")
    from .spreadsheet_formula_checks import validate_syntax

    if strict:
        validate_syntax(tokens)
    for index, token in enumerate(tokens):
        if strict and token.type == "FUNC" and token.subtype == "OPEN":
            name = token.value[:-1].upper()
            arguments, _ = _function_arguments(tokens, index)
            check_function_ranges(name, arguments, current, dimensions)
        if (
            token.type == "FUNC"
            and token.subtype == "OPEN"
            and token.value[:-1].upper() == "XLOOKUP"
        ):
            arguments, _ = _function_arguments(tokens, index)
            if len(arguments) not in {3, 4} or any(not argument for argument in arguments):
                raise ValueError("XLOOKUP unterstützt genaues Suchen mit drei oder vier Argumenten")
            shapes = []
            for argument in arguments[1:3]:
                reference = "".join(item.value for item in argument).strip()
                _, left, top, right, bottom = bounded_reference(reference, current, dimensions)
                width, height = right - left + 1, bottom - top + 1
                if width > 1 and height > 1:
                    raise ValueError("XLOOKUP benötigt eindimensionale Zellbereiche")
                shapes.append((width, height))
            if shapes[0] != shapes[1]:
                raise ValueError("XLOOKUP-Such- und Ergebnisbereich müssen gleich groß sein")
    stack: list[str] = []
    references: set[tuple[str, int, int]] = set()
    previous = ""
    for token in tokens:
        kind, sub, value = token.type, token.subtype, token.value
        if kind == "WHITE-SPACE":
            continue
        if kind in {"FUNC", "PAREN"}:
            if sub == "OPEN":
                if kind == "FUNC" and value[:-1].upper() not in FUNCTIONS:
                    raise ValueError("Excel-Funktion nicht freigegeben: " + value[:-1])
                stack.append(kind)
            elif not stack or stack.pop() != kind:
                raise ValueError("Formelklammern sind ungültig")
        elif kind == "OPERAND":
            if previous in {"operand", "close"}:
                raise ValueError("Implizite Verweise sind nicht erlaubt")
            if sub == "RANGE":
                sheet, left, top, right, bottom = bounded_reference(value, current, dimensions)
                if len(references) + (right - left + 1) * (bottom - top + 1) > 20_000:
                    raise ValueError("Formelbereich überschreitet das Abhängigkeitslimit")
                references.update(
                    (sheet, column, row)
                    for column in range(left, right + 1)
                    for row in range(top, bottom + 1)
                )
            elif sub not in {"NUMBER", "TEXT", "LOGICAL"}:
                raise ValueError("Ungültiger Formelausdruck")
            elif strict and sub == "NUMBER" and not math.isfinite(float(value)):
                raise ValueError("Formelzahl muss endlich sein")
        elif kind == "SEP":
            if sub != "ARG" or value != "," or "FUNC" not in stack:
                raise ValueError("Nur englische Formeln mit Komma als Argumenttrenner erlaubt")
        elif kind == "OPERATOR-INFIX":
            if value not in {"+", "-", "*", "/", "^", "&", "=", "<>", "<", ">", "<=", ">="}:
                raise ValueError("Formeloperator nicht unterstützt")
        elif kind == "OPERATOR-PREFIX":
            if value not in {"+", "-"}:
                raise ValueError("Formeloperator nicht unterstützt")
        elif kind != "OPERATOR-POSTFIX" or value != "%":
            raise ValueError("Formelausdruck nicht unterstützt")
        previous = "operand" if kind == "OPERAND" else "close" if sub == "CLOSE" else "other"
    if stack:
        raise ValueError("Formelklammern sind unvollständig")
    calculation_formula(formula)
    return references


def check_function_ranges(
    name: str, arguments: list[list[Any]], current: str, dimensions: dict[str, tuple[int, int]]
) -> None:
    def shape(argument: list[Any]) -> tuple[int, int] | None:
        values = [item for item in argument if item.type != "WHITE-SPACE"]
        if len(values) != 1 or values[0].type != "OPERAND" or values[0].subtype != "RANGE":
            return None
        _, left, top, right, bottom = bounded_reference(values[0].value, current, dimensions)
        return right - left + 1, bottom - top + 1

    indexes: list[int] = []
    if name in {"SUMIF", "AVERAGEIF"} and len(arguments) == 3:
        indexes = [0, 2]
    elif name in {"SUMIFS", "AVERAGEIFS"}:
        indexes = [0, *range(1, len(arguments), 2)]
    elif name == "COUNTIFS":
        indexes = list(range(0, len(arguments), 2))
    elif name == "SUMPRODUCT":
        indexes = list(range(len(arguments)))
    shapes = [shape(arguments[index]) for index in indexes]
    known = [value for value in shapes if value is not None]
    if known and any(value != known[0] for value in known[1:]):
        raise ValueError("Funktionsbereiche müssen gleich groß sein: " + name)

    if name in {"VLOOKUP", "HLOOKUP"}:
        bounds = shape(arguments[1])
        values = [item for item in arguments[2] if item.type != "WHITE-SPACE"]
        if bounds and len(values) == 1 and values[0].subtype == "NUMBER":
            number = float(values[0].value)
            maximum = bounds[0] if name == "VLOOKUP" else bounds[1]
            if not number.is_integer() or not 1 <= number <= maximum:
                raise ValueError("Suchindex außerhalb des Tabellenbereichs: " + name)
    if name == "SUBTOTAL":
        values = [item for item in arguments[0] if item.type != "WHITE-SPACE"]
        if len(values) == 1 and values[0].subtype == "NUMBER":
            number = float(values[0].value)
            if not number.is_integer() or number not in {*range(1, 12), *range(101, 112)}:
                raise ValueError("SUBTOTAL benötigt eine gültige Funktionsnummer")


def contains_subtotal(formula: str) -> bool:
    return any(
        token.type == "FUNC" and token.subtype == "OPEN" and token.value[:-1].upper() == "SUBTOTAL"
        for token in Tokenizer(formula).items
    )


def aggregate_references(
    formula: str, current: str, dimensions: dict[str, tuple[int, int]]
) -> set[tuple[str, int, int]]:
    """Unambiguous summary formulas, not taxes/derived per-row calculations.

    SUM(B2:B3) and ROUND(SUM(B2:B3),2) summarize the same data. In contrast,
    SUM(B2:B3)*0.2 may be an explicitly requested tax line in a valid total.
    We cannot infer business intent and must not reject that arithmetic.
    """
    tokens = [token for token in Tokenizer(formula).items if token.type != "WHITE-SPACE"]
    result: set[tuple[str, int, int]] = set()
    functions = {
        "SUM",
        "AVERAGE",
        "MIN",
        "MAX",
        "COUNT",
        "COUNTA",
        "MEDIAN",
        "COUNTBLANK",
        "SUMPRODUCT",
        "LARGE",
        "SMALL",
        "NPV",
        "IRR",
        "SUMIF",
        "SUMIFS",
        "AVERAGEIF",
        "AVERAGEIFS",
        "COUNTIF",
        "COUNTIFS",
    }
    while tokens and tokens[0].type == "PAREN" and tokens[-1].type == "PAREN":
        depth = 0
        for _index, token in enumerate(tokens):
            if token.type in {"PAREN", "FUNC"}:
                depth += 1 if token.subtype == "OPEN" else -1
            if depth == 0:
                break
        if _index != len(tokens) - 1:
            break
        tokens = tokens[1:-1]
    if not tokens or tokens[0].type != "FUNC" or tokens[0].subtype != "OPEN":
        return result
    name = tokens[0].value[:-1].upper()
    arguments, end = _function_arguments(tokens, 0)
    if end != len(tokens) - 1:
        return result
    if name in {"ROUND", "ROUNDUP", "ROUNDDOWN"}:
        return aggregate_references(
            "=" + "".join(token.value for token in arguments[0]), current, dimensions
        )
    if name in functions:
        if name in {"SUMIF", "AVERAGEIF"}:
            arguments = [arguments[2] if len(arguments) == 3 else arguments[0]]
        elif name in {"SUMIFS", "AVERAGEIFS"}:
            arguments = arguments[:1]
        elif name == "COUNTIF":
            arguments = arguments[:1]
        elif name == "COUNTIFS":
            arguments = arguments[::2]
        elif name == "NPV":
            arguments = arguments[1:]
        elif name in {"IRR", "LARGE", "SMALL"}:
            arguments = arguments[:1]
        for argument in arguments:
            # A reduction over explicit ranges is a summary. An expression such
            # as SUM(B2*0.2) instead computes a derived detail value.
            if len(argument) != 1 or argument[0].type != "OPERAND":
                return set()
            for value in argument:
                if value.type == "OPERAND" and value.subtype == "RANGE":
                    sheet, left, top, right, bottom = bounded_reference(
                        value.value, current, dimensions
                    )
                    result.update(
                        (sheet, column, row)
                        for column in range(left, right + 1)
                        for row in range(top, bottom + 1)
                    )
    return result


def validate_workbook(plan: SpreadsheetPlan, *, strict: bool = True) -> None:
    from .artifact_specs import FormulaCell

    dimensions = {
        sheet.name: (len(sheet.columns), len(sheet.rows) + 1 + bool(sheet.totals))
        for sheet in plan.sheets
    }
    graph: dict[tuple[str, int, int], set[tuple[str, int, int]]] = {}
    dependency_count = 0
    for sheet in plan.sheets:
        for row_number, row in enumerate(sheet.rows, 2):
            for column, cell in enumerate(row, 1):
                if isinstance(cell, FormulaCell):
                    graph[(sheet.name, column, row_number)] = formula_references(
                        cell.formula, sheet.name, dimensions, strict=strict
                    )
                    dependency_count += len(graph[(sheet.name, column, row_number)])
                    if dependency_count > 200_000:
                        raise ValueError(
                            "Formelabhängigkeiten überschreiten das Arbeitsmappenlimit"
                        )
        for column_name in sheet.number_formats:
            if ord(column_name) - ord("A") + 1 > len(sheet.columns):
                raise ValueError("Zahlenformat verweist auf fehlende Spalte")
        for column_name in sheet.totals:
            if not sheet.rows or ord(column_name) - ord("A") + 1 > len(sheet.columns):
                raise ValueError("Ergebniszeile benötigt Daten und vorhandene Spalten")
            column_number = ord(column_name) - ord("A") + 1
            covered = {
                (sheet.name, column_number, number) for number in range(2, len(sheet.rows) + 2)
            }
            # SUBTOTAL ignores nested SUBTOTAL cells. SUM/AVERAGE/etc summary
            # rows are not ignored, and would count already summarized data twice.
            ignored = set()
            for row_number, row in enumerate(sheet.rows, 2):
                cell = row[column_number - 1]
                if not isinstance(cell, FormulaCell):
                    continue
                node = (sheet.name, column_number, row_number)
                if contains_subtotal(cell.formula):
                    ignored.add(node)
                elif (
                    strict and aggregate_references(cell.formula, sheet.name, dimensions) & covered
                ):
                    raise ValueError(
                        "Doppelte Aggregation: automatische Ergebniszeile in "
                        + sheet.name
                        + "!"
                        + column_name
                        + " enthält bereits aggregierte Daten. Entweder totals oder eine "
                        "manuelle Aggregatzeile verwenden; Detailformeln bleiben erlaubt."
                    )
            graph[(sheet.name, column_number, len(sheet.rows) + 2)] = covered - ignored
        for coordinate in sheet.highlight_cells:
            bounded_reference(coordinate, sheet.name, dimensions)
        for group in sheet.column_groups:
            if group.start > group.end or ord(group.end) - ord("A") >= len(sheet.columns):
                raise ValueError("Spaltengruppe außerhalb der Tabelle")
        for chart in sheet.charts:
            _, left, top, right, bottom = bounded_reference(
                chart.categories, sheet.name, dimensions
            )
            if left != right:
                raise ValueError("Diagrammkategorien müssen eine Spalte verwenden")
            for series in chart.series:
                _, first, start, last, finish = bounded_reference(series, sheet.name, dimensions)
                count = finish - start + 1
                category_count = bottom - top + 1
                if first != last or not (
                    (start == 1 and count == category_count + 1)
                    or (start > 1 and count == category_count)
                ):
                    raise ValueError("Diagrammreihe benötigt eine passende Anzahl Werte")
            if chart.type == "pie" and len(chart.series) != 1:
                raise ValueError("Kreisdiagramm benötigt genau eine Datenreihe")
    if sum(len(references) for references in graph.values()) > 200_000:
        raise ValueError("Formelabhängigkeiten überschreiten das Arbeitsmappenlimit")
    # Iterative traversal avoids recursion limits for large but valid formula chains.
    done: set[tuple[str, int, int]] = set()
    for node in graph:
        visiting: set[tuple[str, int, int]] = set()
        stack = [(node, False)]
        while stack:
            item, exiting = stack.pop()
            if exiting:
                visiting.remove(item)
                done.add(item)
            elif item in visiting:
                raise ValueError("Zirkelbezug in Excel-Formeln")
            elif item not in done:
                visiting.add(item)
                stack.append((item, True))
                stack.extend(
                    (dependency, False) for dependency in graph[item] if dependency in graph
                )


def decorate_sheet(sheet: Any, item: Any, index: int) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.row_dimensions[1].height = 31.5
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            color = (
                "008000"
                if cell.data_type == "f" and "!" in cell.value
                else "000000"
                if cell.data_type == "f"
                else "0000FF"
            )
            cell.font = Font(name="Liberation Sans", size=11, color=color)
            cell.number_format = (
                FORMATS.get(
                    item.number_formats.get(cell.column_letter, "integer"), FORMATS["integer"]
                )
                if cell.data_type == "f" or isinstance(cell.value, (int, float))
                else "General"
            )
    if item.rows and len({name.casefold() for name in item.columns}) == len(item.columns):
        table = Table(displayName=f"tbl_data_{index}", ref=sheet.dimensions)
        table.tableColumns = [
            TableColumn(id=number, name=name, totalsRowFunction=item.totals.get(chr(64 + number)))
            for number, name in enumerate(item.columns, 1)
        ]
        if item.totals:
            table.totalsRowCount = 1
            table.totalsRowShown = True
            table.autoFilter = AutoFilter(
                ref=f"A1:{sheet.cell(len(item.rows) + 1, len(item.columns)).coordinate}"
            )
            sheet.auto_filter.ref = table.autoFilter.ref
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleLight8", showRowStripes=not item.no_stripes, showColumnStripes=False
        )
        sheet.add_table(table)
    for coordinate in item.highlight_cells:
        sheet[coordinate].fill = PatternFill("solid", fgColor="FFF2CC")
    for group in item.column_groups:
        sheet.column_dimensions.group(group.start, group.end, hidden=True)
    constructors = {
        "bar": BarChart,
        "line": LineChart,
        "area": AreaChart,
        "pie": PieChart,
        "scatter": ScatterChart,
    }
    palette = ["E6531F", "020623", "2962FF", "008080", "E6A23C"]
    for specification in item.charts:
        chart = constructors[specification.type]()
        chart.title = specification.title
        chart.style = 10
        chart.width, chart.height = 19, 10
        left, top, right, bottom = range_boundaries(specification.categories)
        categories = Reference(sheet, min_col=left, min_row=top, max_col=right, max_row=bottom)
        for data_range in specification.series:
            left, top, right, bottom = range_boundaries(data_range)
            values = Reference(sheet, min_col=left, min_row=top, max_col=right, max_row=bottom)
            if specification.type == "scatter":
                chart.series.append(Series(values, categories, title_from_data=top == 1))
            else:
                chart.add_data(values, titles_from_data=top == 1)
            if top > 1:
                header = Reference(sheet, min_col=left, min_row=1, max_col=left, max_row=1)
                chart.series[-1].tx = SeriesLabel(strRef=StrRef(f=str(header)))
        if specification.type != "scatter":
            chart.set_categories(categories)
        for number, series in enumerate(chart.series):
            series.graphicalProperties.solidFill = palette[number % len(palette)]
            series.graphicalProperties.line.solidFill = palette[number % len(palette)]
        sheet.add_chart(chart, specification.anchor)


def cache_formula_results(path: Any) -> None:
    """Calculate with the fixed offline Calc binary; inject caches into original ZIP.

    Calc's export is only an intermediate calculation result. All original table
    styles, formula text, charts and relationships remain byte-identical.
    """
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    tag = "{" + namespace + "}"
    original = load_workbook(path, read_only=True, data_only=False)
    formula_cells = {
        (sheet.title, cell.coordinate)
        for sheet in original
        for row in sheet
        for cell in row
        if cell.data_type == "f"
    }
    names = original.sheetnames
    original.close()
    if not formula_cells:
        return
    with tempfile.TemporaryDirectory(prefix="calc-", dir=path.parent) as temporary:
        folder = Path(temporary)
        profile = folder / "profile"
        converted = folder / path.name
        input_folder = folder / "input"
        input_folder.mkdir()
        calculation_input = input_folder / path.name
        with (
            zipfile.ZipFile(path) as source,
            zipfile.ZipFile(calculation_input, "w", zipfile.ZIP_DEFLATED) as output,
        ):
            for entry in source.infolist():
                data = source.read(entry.filename)
                if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", entry.filename):
                    root = ET.fromstring(data)
                    for formula in root.iter(tag + "f"):
                        if formula.text:
                            formula.text = calculation_formula("=" + formula.text)[1:]
                    data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                output.writestr(entry, data)
        command = [
            "/usr/lib/libreoffice/program/soffice.bin" if sys.platform == "linux" else "soffice",
            "-env:UserInstallation=" + profile.as_uri(),
            "-env:OSL_SOCKET_PATH=/run/ipc",
            "--headless",
            "--norestore",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(folder),
            str(calculation_input),
        ]
        process = subprocess.run(command, capture_output=True, timeout=60)
        if sys.platform == "linux" and process.returncode == 81:
            process = subprocess.run(command, capture_output=True, timeout=60)
        process.check_returncode()
        if (
            not converted.is_file()
            or converted.is_symlink()
            or converted.stat().st_size > 40 * 1024 * 1024
        ):
            raise ValueError("Excel-Neuberechnung fehlt oder überschreitet das Limit")
        cached = load_workbook(converted, read_only=True, data_only=True)
        empty_strings: set[tuple[str, str]] = set()
        with zipfile.ZipFile(converted) as calculation:
            for number, name in enumerate(names, 1):
                root = ET.fromstring(calculation.read(f"xl/worksheets/sheet{number}.xml"))
                for cell in root.iter(tag + "c"):
                    value_node = cell.find(tag + "v")
                    if (
                        cell.attrib.get("t") == "str"
                        and value_node is not None
                        and not value_node.text
                    ):
                        empty_strings.add((name, cell.attrib.get("r", "")))
        values: dict[tuple[str, str], Any] = {}
        for name, coordinate in formula_cells:
            cell = cached[name][coordinate]
            if cell.data_type == "e" or (
                cell.value is None and (name, coordinate) not in empty_strings
            ):
                cached.close()
                raise ValueError("Excel-Formel nicht berechenbar: " + name + "!" + coordinate)
            value = "" if (name, coordinate) in empty_strings else cell.value
            if isinstance(value, (date, datetime)):
                value = to_excel(value, cached.epoch)
            elif isinstance(value, time):
                value = time_to_days(value)
            elif isinstance(value, timedelta):
                value = timedelta_to_days(value)
            if isinstance(value, float) and not math.isfinite(value):
                cached.close()
                raise ValueError("Excel-Formel liefert keine endliche Zahl")
            values[(name, coordinate)] = value
        cached.close()
        rewritten = folder / "cached.xlsx"
        with (
            zipfile.ZipFile(path) as source,
            zipfile.ZipFile(rewritten, "w", zipfile.ZIP_DEFLATED) as output,
        ):
            for entry in source.infolist():
                data = source.read(entry.filename)
                match = re.fullmatch(r"xl/worksheets/sheet(\d+)\.xml", entry.filename)
                if match:
                    name = names[int(match[1]) - 1]
                    root = ET.fromstring(data)
                    for cell in root.iter(tag + "c"):
                        key = (name, cell.attrib.get("r", ""))
                        if key not in values:
                            continue
                        for child in list(cell):
                            if child.tag in {tag + "v", tag + "is"}:
                                cell.remove(child)
                        value = values[key]
                        cell.attrib["t"] = (
                            "b"
                            if isinstance(value, bool)
                            else "n"
                            if isinstance(value, (int, float))
                            else "str"
                        )
                        ET.SubElement(cell, tag + "v").text = (
                            str(int(value)) if isinstance(value, bool) else str(value)
                        )
                    data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
                output.writestr(entry, data)
        rewritten.replace(path)
