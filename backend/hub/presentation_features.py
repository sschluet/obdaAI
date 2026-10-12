"""Fixed native Office diagram adapter; uploaded scripts are never executed."""

import math
import sys
from functools import lru_cache
from pathlib import Path
from typing import Any

from PIL import ImageFont
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from hub.artifact_specs import Slide, SlideVisual, VisualItem

NAVY = "020623"
ORANGE = "EB601D"
BLUE = "4778E2"
INK2 = "4A5062"
RULE = "D9DBE3"
PAPER = "F4F5F8"
STATUS = [ORANGE, NAVY, "8E92A3", "D0D2DB"]
SERIES = [ORANGE, BLUE, NAVY, "8E92A3"]


def foreground(fill: str) -> str:
    def luminance(hex_color: str) -> float:
        values = [int(hex_color[i : i + 2], 16) / 255 for i in (0, 2, 4)]
        linear = [v / 12.92 if v <= 0.04045 else ((v + 0.055) / 1.055) ** 2.4 for v in values]
        return sum(v * w for v, w in zip(linear, (0.2126, 0.7152, 0.0722), strict=True))

    lum = luminance(fill)
    white = 1.05 / (lum + 0.05)
    navy = (lum + 0.05) / (luminance(NAVY) + 0.05)
    return "FFFFFF" if white > navy else NAVY


@lru_cache(maxsize=100)
def _font(size: float, bold: bool) -> Any:
    suffix = "Bold" if bold else "Regular"
    candidates = (
        [f"/System/Library/Fonts/Supplemental/Arial{' Bold' if bold else ''}.ttf"]
        if sys.platform == "darwin"
        else [
            f"/usr/share/fonts/truetype/liberation/LiberationSans-{suffix}.ttf",
            f"/usr/share/fonts/truetype/liberation2/LiberationSans-{suffix}.ttf",
        ]
    )
    path = next((path for path in candidates if Path(path).is_file()), None)
    if path is None:
        raise ValueError("Schrift für die Schaubildprüfung fehlt")
    return ImageFont.truetype(path, round(size * 96 / 72))


def text(
    slide: Any,
    x: float,
    y: float,
    w: float,
    h: float,
    value: str,
    size: float = 12,
    color: str = NAVY,
    bold: bool = False,
    center: bool = False,
) -> Any:
    # Bound wrapping before writing native text. Shrink only within a readable range;
    # oversize authored content is rejected instead of hidden below another element.
    minimum = min(size, 8.5)
    chosen_size = size
    while True:
        font = _font(chosen_size, bold)
        count = 0
        for part in value.split("\n"):
            current = ""
            count += 1
            for character in part:
                if current and font.getlength(current + character) > w * 96:
                    count += 1
                    current = ""
                current += character
        required = count * chosen_size * 1.13
        if required <= h * 72 or not value:
            break
        chosen_size -= 0.5
        if chosen_size < minimum:
            raise ValueError("Schaubildtext passt nicht in die Fläche; Text kürzen")
    size = chosen_size
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    for n, line in enumerate(value.split("\n")):
        p = frame.paragraphs[0] if n == 0 else frame.add_paragraph()
        p.text = line
        p.font.name = "Arial"
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)  # type: ignore[no-untyped-call]
        p.space_after = Pt(0)
        p.line_spacing = 1.13
        p.alignment = PP_ALIGN.CENTER if center else PP_ALIGN.LEFT
    return box


def shape(
    slide: Any, x: float, y: float, w: float, h: float, color: str, kind: Any = MSO_SHAPE.RECTANGLE
) -> Any:
    obj = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    obj.fill.solid()
    obj.fill.fore_color.rgb = RGBColor.from_string(color)  # type: ignore[no-untyped-call]
    obj.line.fill.background()
    return obj


def line(
    slide: Any, x1: float, y1: float, x2: float, y2: float, color: str = RULE, width: float = 0.8
) -> None:
    obj = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    obj.line.color.rgb = RGBColor.from_string(color)  # type: ignore[no-untyped-call]
    obj.line.width = Pt(width)


def marker(slide: Any, x: float, y: float, value: str, fill: str = ORANGE) -> None:
    obj = shape(slide, x, y, 0.35, 0.35, fill, MSO_SHAPE.OVAL)
    obj.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = obj.text_frame.paragraphs[0]
    p.text = value
    p.alignment = PP_ALIGN.CENTER
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = RGBColor.from_string(foreground(fill))  # type: ignore[no-untyped-call]
    obj.text_frame.margin_left = obj.text_frame.margin_right = 0
    obj.text_frame.margin_top = obj.text_frame.margin_bottom = 0


def native_chart(slide: Any, visual: SlideVisual, area: tuple[float, float, float, float]) -> None:
    data = CategoryChartData()  # type: ignore[no-untyped-call]
    data.categories = visual.columns
    for series in visual.series:
        data.add_series(series.name, series.values)  # type: ignore[no-untyped-call]
    kinds = {
        "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
        "bar": XL_CHART_TYPE.BAR_CLUSTERED,
        "line": XL_CHART_TYPE.LINE,
        "area": XL_CHART_TYPE.AREA,
        "pie": XL_CHART_TYPE.PIE,
        "doughnut": XL_CHART_TYPE.DOUGHNUT,
    }
    chart = slide.shapes.add_chart(kinds[visual.chart_type], *(Inches(v) for v in area), data).chart
    chart.has_legend = len(visual.series) > 1 or visual.chart_type in {"pie", "doughnut"}
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.font.name = "Arial"
        chart.legend.font.size = Pt(10)
    chart.chart_style = 10
    for index, series in enumerate(chart.series):
        series.format.fill.solid()
        series.format.fill.fore_color.rgb = RGBColor.from_string(SERIES[index])  # type: ignore[no-untyped-call]
        series.format.line.color.rgb = RGBColor.from_string(SERIES[index])  # type: ignore[no-untyped-call]
        if visual.chart_type in {"pie", "doughnut"}:
            for j, point in enumerate(series.points):
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = RGBColor.from_string(STATUS[j % 4])  # type: ignore[no-untyped-call]
    if visual.chart_type not in {"pie", "doughnut"}:
        for axis in (chart.category_axis, chart.value_axis):
            axis.tick_labels.font.name = "Arial"
            axis.tick_labels.font.size = Pt(10)
            axis.tick_labels.font.color.rgb = RGBColor.from_string(INK2)  # type: ignore[no-untyped-call]


def _item(
    slide: Any, item: VisualItem, index: int, area: tuple[float, float, float, float]
) -> None:
    x, y, w, h = area
    compact = h < 1.15
    label_h = 0.28 if compact else 0.45
    text(slide, x, y, 0.65, min(0.55, h), f"{index + 1:02}", 24 if compact else 28, ORANGE)
    text(slide, x + 0.8, y, w - 0.8, label_h, item.label, 11 if compact else 15, bold=True)
    text(
        slide,
        x + 0.8,
        y + label_h + 0.07,
        w - 0.8,
        max(0.2, h - label_h - 0.14),
        item.text,
        9 if compact else 12,
        INK2,
    )


def period_bounds(plan: SlideVisual, item: VisualItem) -> tuple[float, float]:
    if plan.kind == "matrix":
        start = float(item.period or 0)
        return start, start + 1
    if item.from_period is not None and item.to_period is not None:
        return float(plan.columns.index(item.from_period)), float(
            plan.columns.index(item.to_period) + 1
        )
    return float(item.start or 0), float(item.end or 0)


def timeline_rows(plan: SlideVisual) -> list[tuple[str, list[VisualItem]]]:
    """Pack nonoverlapping intervals, never paint one task over another."""
    result: list[tuple[str, list[VisualItem]]] = []
    for lane in dict.fromkeys(i.lane for i in plan.items):
        rows: list[list[VisualItem]] = []
        ends: list[float] = []
        for item in sorted(
            (i for i in plan.items if i.lane == lane), key=lambda i: period_bounds(plan, i)
        ):
            start, end = period_bounds(plan, item)
            # Milestone labels span a row, so these reserve a complete subrow.
            if end - start < 0.1:
                start, end = 0, float(len(plan.columns))
            index = next((n for n, previous in enumerate(ends) if previous <= start), len(rows))
            if index == len(rows):
                rows.append([])
                ends.append(end)
            else:
                ends[index] = end
            rows[index].append(item)
        result.extend((lane, members) for members in rows)
    if len(result) > 6:
        raise ValueError(
            "Zeitdarstellung zu dicht; maximal sechs überlappungsfreie Zeilen je Folie"
        )
    return result


def visual(slide: Any, plan: SlideVisual, area: tuple[float, float, float, float]) -> None:
    x, y, w, h = area
    items = plan.items
    if plan.kind == "chart":
        native_chart(slide, plan, area)
    elif plan.kind == "process":
        gap = 0.38
        cw = (w - gap * (len(items) - 1)) / len(items)
        for i, item in enumerate(items):
            px = x + i * (cw + gap)
            shape(slide, px, y + 0.32, cw, 1, NAVY)
            text(slide, px + 0.1, y + 0.55, cw - 0.2, 0.65, item.label, 13, "FFFFFF", True, True)
            marker(slide, px - 0.06, y + 0.15, str(i + 1))
            text(slide, px, y + 1.65, cw, min(2.8, h - 1.65), item.text, 12, INK2)
            if i < len(items) - 1:
                shape(slide, px + cw + 0.07, y + 0.63, 0.25, 0.45, ORANGE, MSO_SHAPE.CHEVRON)
    elif plan.kind == "swimlane":
        lanes = list(dict.fromkeys(i.lane for i in items))
        rh = h / len(lanes)
        for li, lane in enumerate(lanes):
            members = [i for i in items if i.lane == lane]
            yy = y + li * rh
            text(slide, x, yy + 0.2, 1.7, 0.7, lane, 12, bold=True)
            cw = (w - 1.9 - 0.23 * (len(members) - 1)) / len(members)
            for j, item in enumerate(members):
                xx = x + 1.9 + j * (cw + 0.23)
                shape(slide, xx, yy + 0.12, cw, 0.65, NAVY)
                text(
                    slide,
                    xx + 0.07,
                    yy + 0.22,
                    cw - 0.14,
                    0.45,
                    item.label,
                    11,
                    "FFFFFF",
                    True,
                    True,
                )
                if j < len(members) - 1:
                    shape(slide, xx + cw + 0.045, yy + 0.35, 0.14, 0.16, ORANGE, MSO_SHAPE.CHEVRON)
                text(slide, xx, yy + 0.87, cw, max(0.28, rh - 1), item.text, 10, INK2)
            if li < len(lanes) - 1:
                line(slide, x, yy + rh - 0.05, x + w, yy + rh - 0.05)
    elif plan.kind == "cycle":
        diameter = min(4, w * 0.46, h - 0.25)
        cx, cy = x + diameter / 2, y + h / 2
        radius = diameter * 0.35
        for i, item in enumerate(items):
            angle = i * 2 * math.pi / len(items) - math.pi / 2
            next_a = (i + 1) * 2 * math.pi / len(items) - math.pi / 2
            fill = STATUS[min(i, 3)]
            arc = shape(
                slide, cx - radius, cy - radius, 2 * radius, 2 * radius, fill, MSO_SHAPE.BLOCK_ARC
            )
            arc.adjustments[0] = (math.degrees(angle) + 4) % 360
            arc.adjustments[1] = (math.degrees(next_a) - 5) % 360
            arc.adjustments[2] = 0.12
            tip_x = cx + radius * 0.88 * math.cos(next_a - 0.08)
            tip_y = cy + radius * 0.88 * math.sin(next_a - 0.08)
            shape(
                slide, tip_x - 0.14, tip_y - 0.14, 0.28, 0.28, fill, MSO_SHAPE.ISOSCELES_TRIANGLE
            ).rotation = math.degrees(next_a) + 180
            mid_angle = (angle + next_a) / 2
            marker(
                slide,
                cx + radius * 0.85 * math.cos(mid_angle) - 0.175,
                cy + radius * 0.85 * math.sin(mid_angle) - 0.175,
                str(i + 1),
                fill,
            )
            _item(
                slide,
                item,
                i,
                (x + diameter + 0.3, y + i * h / len(items), w - diameter - 0.3, h / len(items)),
            )
        text(slide, cx - 0.65, cy - 0.45, 1.3, 0.6, plan.center_value, 30, center=True)
        text(slide, cx - 0.7, cy + 0.18, 1.4, 0.6, plan.center_label, 10, center=True)
    elif plan.kind in {"timeline", "matrix"}:
        packed_rows = timeline_rows(plan)
        label_w = min(2.3, w * 0.23)
        cw = (w - label_w) / len(plan.columns)
        rh = (h - 0.7) / len(packed_rows)
        for c, column in enumerate(plan.columns):
            xx = x + label_w + c * cw
            text(slide, xx, y, cw, 0.4, column, 10, bold=True, center=True)
            line(slide, xx, y + 0.55, xx, y + h)
        for n, (lane, members) in enumerate(packed_rows):
            yy = y + 0.7 + n * rh
            text(slide, x, yy, label_w - 0.15, rh, lane, 11, bold=True)
            line(slide, x, yy + rh, x + w, yy + rh)
            for j, item in enumerate(members):
                start, end = period_bounds(plan, item)
                fill = STATUS[j % 4]
                xx = x + label_w + start * cw + 0.05
                width = (end - start) * cw - 0.1
                if width < 0.18:
                    marker_x = min(x + w - 0.25, max(x + label_w, xx))
                    shape(slide, marker_x, yy + 0.16, 0.18, 0.18, fill, MSO_SHAPE.DIAMOND)
                    text(
                        slide,
                        x + label_w + 0.1,
                        yy + 0.4,
                        w - label_w - 0.2,
                        max(0.2, rh - 0.4),
                        item.label,
                        9,
                    )
                    continue
                shape(slide, xx, yy + 0.12, width, min(0.65, rh - 0.2), fill)
                text(
                    slide,
                    xx + 0.06,
                    yy + 0.22,
                    max(0.1, width - 0.12),
                    min(0.45, rh - 0.3),
                    item.label,
                    9,
                    foreground(fill),
                    True,
                )
    elif plan.kind == "kpi":
        cw = w / len(items)
        for i, item in enumerate(items):
            xx = x + cw * i
            text(slide, xx, y + 0.35, cw - 0.35, 1, item.value, 50)
            text(slide, xx, y + 1.65, cw - 0.35, 0.8, item.label, 15, bold=True)
            text(slide, xx, y + 2.65, cw - 0.35, h - 2.65, item.text, 12, INK2)
            if i:
                line(slide, xx - 0.18, y, xx - 0.18, y + h)
    elif plan.kind == "layers":
        rh = h / len(items)
        for i, item in enumerate(items):
            yy = y + i * rh
            fill = STATUS[min(i, 3)]
            shape(slide, x, yy, w * 0.36, rh - 0.12, fill)
            text(
                slide,
                x + 0.15,
                yy + 0.14,
                w * 0.36 - 0.3,
                rh - 0.25,
                item.label,
                12,
                foreground(fill),
                True,
            )
            text(slide, x + w * 0.41, yy + 0.14, w * 0.59, rh - 0.25, item.text, 12, INK2)
    elif plan.kind == "comparison":
        lanes = list(dict.fromkeys(i.lane for i in items))
        cw = (w - 0.95) / 2
        for n, lane in enumerate(lanes):
            xx = x + n * (cw + 0.95)
            fill = PAPER if n == 0 else NAVY
            shape(slide, xx, y, cw, h, fill)
            text(slide, xx + 0.2, y + 0.2, cw - 0.4, 0.4, lane.upper(), 12, foreground(fill), True)
            members = [i for i in items if i.lane == lane]
            rh = (h - 0.95) / len(members)
            for i, item in enumerate(members):
                text(
                    slide,
                    xx + 0.2,
                    y + 0.95 + i * rh,
                    cw - 0.4,
                    rh,
                    item.label + ("\n" + item.text if item.text else ""),
                    12,
                    foreground(fill),
                )
        shape(slide, x + cw + 0.12, y + h / 2 - 0.3, 0.7, 0.6, ORANGE, MSO_SHAPE.RIGHT_ARROW)
    elif plan.kind == "dot_matrix":
        total = sum(int(i.value) for i in items)
        cols = min(15, math.ceil(math.sqrt(total * w / max(h - 1, 1))))
        rows = math.ceil(total / cols)
        d = min(0.3, (w - 0.1) / cols, (h - 1) / rows) * 0.7
        index = 0
        for j, item in enumerate(items):
            for _ in range(int(item.value)):
                shape(
                    slide,
                    x + index % cols * (d / 0.7),
                    y + index // cols * (d / 0.7),
                    d,
                    d,
                    STATUS[j],
                    MSO_SHAPE.OVAL,
                )
                index += 1
            text(
                slide,
                x + j * w / len(items),
                y + h - 0.8,
                w / len(items) - 0.1,
                0.7,
                f"{item.value} · {item.label}",
                10,
            )
    elif plan.kind == "status_bars":
        totals = [sum(s.values[i] for s in plan.series) for i in range(len(plan.columns))]
        maximum = max(totals) or 1
        rh = (h - 0.5) / len(plan.columns)
        for row, category in enumerate(plan.columns):
            yy = y + row * rh
            text(slide, x, yy + 0.08, 2, 0.4, category, 11)
            offset = x + 2.2
            for j, series in enumerate(plan.series):
                width = (w - 2.85) * series.values[row] / maximum
                if width > 0:
                    shape(slide, offset, yy, width, min(0.45, rh - 0.12), STATUS[j])
                    if width > 0.35:
                        text(
                            slide,
                            offset + 0.04,
                            yy + 0.07,
                            width - 0.08,
                            0.3,
                            f"{series.values[row]:g}",
                            9,
                            foreground(STATUS[j]),
                            center=True,
                        )
                    offset += width
            text(slide, offset + 0.1, yy + 0.07, 0.6, 0.3, f"{totals[row]:g}", 10)
        text(slide, x, y + h - 0.3, w, 0.3, " · ".join(s.name for s in plan.series), 9, INK2)
    else:
        rh = h / len(items)
        for i, item in enumerate(items):
            _item(slide, item, i, (x, y + i * rh, w, rh))
            if i < len(items) - 1:
                line(slide, x, y + (i + 1) * rh - 0.12, x + w, y + (i + 1) * rh - 0.12)


def content(slide: Any, item: Slide, number: int, title: str, total: int) -> None:
    # Literal grid and Mac-compatible Arial profile from the reviewed modern guide.
    text(
        slide, 0.52, 0.50, 8, 0.18, (item.kicker or f"{number:02} · {title}").upper(), 8, bold=True
    )
    shape(slide, 0.52, 0.69, 0.32, 0.025, ORANGE)
    size = 20 if len(item.title) >= 69 else 22 if len(item.title) >= 59 else 25
    text(slide, 0.52, 0.74, 10.4, 0.58, item.title, size, bold=True)
    if item.subtitle:
        text(slide, 0.52, 1.36, 11.5, 0.33, item.subtitle, 11, INK2)
    if item.visual:
        area = (0.52, 1.85, 8.7 if item.visual.annotation else 12.29, 4.7)
        visual(slide, item.visual, area)
        if item.visual.annotation:
            shape(slide, 9.55, 1.80, 0.32, 0.025, BLUE)
            text(slide, 9.55, 1.98, 3.26, 4.4, item.visual.annotation, 12, INK2)
    else:
        # Nonvisual legacy plans remain usable, without invented metrics or decorative cards.
        for i, bullet in enumerate(item.bullets):
            _item(
                slide,
                VisualItem(label=bullet[:80], text=bullet if len(bullet) > 80 else ""),
                i,
                (0.52, 1.95 + i * 4.6 / len(item.bullets), 12.29, 4.6 / len(item.bullets)),
            )
    text(slide, 0.52, 6.98, 11, 0.2, item.footer or title, 7.5, INK2)
    text(slide, 12.1, 6.98, 0.7, 0.2, f"{number}/{total}", 7.5, INK2, center=True)
    # An imported generic master can have a different aspect ratio. Keep all native
    # content within its canvas; the corporate master uses the unscaled reference grid.
    sx = slide.part.package.presentation_part.presentation.slide_width / Inches(13.333)
    sy = slide.part.package.presentation_part.presentation.slide_height / Inches(7.5)
    if abs(sx - 1) > 0.001 or abs(sy - 1) > 0.001:
        for obj in slide.shapes:
            obj.left, obj.top = int(obj.left * sx), int(obj.top * sy)
            obj.width, obj.height = int(obj.width * sx), int(obj.height * sy)
    slide.notes_slide.notes_text_frame.text = item.notes
