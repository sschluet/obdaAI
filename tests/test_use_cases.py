"""Declarative forms are local, validated, revision-bound and safe to export."""

import csv
import io
from uuid import uuid4

import pytest
from hub.use_cases import Definition, validated_values
from openpyxl import load_workbook
from pydantic import ValidationError
from use_case_legacy_fixture import legacy

pytest_plugins = ["test_knowledge"]


def spec(kind="table"):
    return {
        "name": "Synthetic monthly overview",
        "fields": [
            {"key": "name", "label": "Object", "type": "text"},
            {"key": "amount", "label": "Amount", "type": "number"},
            {"key": "date", "label": "Date", "type": "date", "required": False},
        ],
        "output_kind": kind,
        "output_columns": [
            {"label": "Amount", "source": "amount"},
            {"label": "Object", "source": "name"},
        ],
        "text_template": "Object {name}: {amount}",
    }


def setup(knowledge, kind="table"):
    app, client, headers, cfg, vault = knowledge
    project = client.post("/api/projects", json={"name": "Synthetic forms"}, headers=headers).json()
    base = f"/api/projects/{project['id']}/use-cases"
    created = legacy(app, spec(kind), project["id"])
    assert created.status_code == 200, created.text
    return client, headers, project, base, created.json()


def submit(client, headers, base, case, name="Synthetic", amount=0, request_id=None):
    return client.post(
        base + f"/{case['id']}/runs",
        json={
            "request_id": request_id or str(uuid4()),
            "expected_revision": case["revision"],
            "values": {"name": name, "amount": amount},
        },
        headers=headers,
    )


def test_form_revision_snapshots_and_exact_replay(knowledge):
    client, headers, project, base, case = setup(knowledge, "chat")
    request_id = str(uuid4())
    first = submit(client, headers, base, case, amount=0, request_id=request_id)
    assert first.status_code == 200
    assert first.json()["output"]["text"] == "Object Synthetic: 0"
    assert first.json()["processing"] == "local_form"
    replay = submit(client, headers, base, case, request_id=request_id)
    assert replay.json()["id"] == first.json()["id"]
    assert submit(client, headers, base, case, amount=10, request_id=request_id).status_code == 409
    updated = client.put(
        base + "/" + case["id"],
        json={
            "spec": {
                **spec("bi"),
                "execution_kind": "ai",
                "processing_prompt": "Synthetic processing",
            },
            "expected_revision": 1,
        },
        headers=headers,
    )
    assert updated.status_code == 200
    assert updated.json()["revision"] == 2
    assert submit(client, headers, base, case).status_code == 409
    assert (
        submit(client, headers, base, case, request_id=request_id).json()["id"]
        == first.json()["id"]
    )
    runs = client.get(base + "/" + case["id"] + "/runs").json()
    assert runs[0]["spec"]["output_kind"] == "chat"
    assert runs[0]["revision"] == 1
    csv_output = client.get(base + "/" + case["id"] + "/output.csv")
    assert (
        len(list(csv.reader(io.StringIO(csv_output.content.decode("utf-8-sig")), delimiter=";")))
        == 1
    )


def test_csv_xlsx_formula_injection_and_numeric_types(knowledge):
    client, headers, project, base, case = setup(knowledge)
    assert (
        submit(
            client, headers, base, case, name='=HYPERLINK("https://example.invalid")', amount=-5
        ).status_code
        == 200
    )
    assert submit(client, headers, base, case, name="  @SUM(1)", amount=3.5).status_code == 200
    rows = list(
        csv.reader(
            io.StringIO(
                client.get(base + "/" + case["id"] + "/output.csv").content.decode("utf-8-sig")
            ),
            delimiter=";",
        )
    )
    assert rows[1][0] == "-5"
    assert rows[1][1].startswith("'=")
    assert rows[2][1].startswith("'  @")
    workbook = load_workbook(
        io.BytesIO(client.get(base + "/" + case["id"] + "/output.xlsx").content)
    )
    assert workbook.active["A2"].value == -5
    assert workbook.active["A2"].data_type == "n"
    assert workbook.active["B2"].data_type == "s"
    assert workbook.active["B2"].hyperlink is None


@pytest.mark.parametrize(
    "values",
    [
        {"name": "A", "amount": True},
        {"name": "A", "amount": "1"},
        {"name": "A", "amount": 1e16},
        {"amount": 1},
        {"name": "A", "amount": 1, "date": "2026-02-30"},
        {"name": "A", "amount": 1, "secret": "x"},
    ],
)
def test_invalid_inputs_rejected(values):
    with pytest.raises(Exception) as error:
        validated_values(Definition.model_validate(spec()), values)
    assert error.value.status_code == 422


def test_invalid_definition_references_and_no_executable_templates():
    for changes in (
        {"output_columns": [{"source": "missing", "label": "x"}]},
        {"text_template": "{name.__class__}"},
        {"fields": [{"key": "name", "label": "X", "type": "choice", "choices": ["A", "A"]}]},
    ):
        with pytest.raises(ValidationError):
            Definition.model_validate({**spec(), **changes})


def test_scope_auth_archive_delete_and_csrf(knowledge):
    client, headers, project, base, case = setup(knowledge)
    second = client.post("/api/projects", json={"name": "Other"}, headers=headers).json()
    assert (
        client.get(f"/api/projects/{second['id']}/use-cases/{case['id']}/runs").status_code == 404
    )
    assert client.post(base, json={"spec": spec()}).status_code == 403
    assert submit(client, headers, base, case).status_code == 200
    assert (
        client.request(
            "DELETE", base + "/" + case["id"], json={"expected_revision": 2}, headers=headers
        ).status_code
        == 409
    )
    assert (
        client.request(
            "DELETE", base + "/" + case["id"], json={"expected_revision": 1}, headers=headers
        ).status_code
        == 200
    )
    assert client.get(base).json() == []
    assert client.get(base + "/" + case["id"] + "/output.xlsx").status_code == 404
    client.patch(
        "/api/projects/" + project["id"],
        json={"name": "Synthetic forms", "revision": 1, "archived": True},
        headers=headers,
    )
    assert client.post(base, json={"spec": spec()}, headers=headers).status_code == 409
    client.cookies.clear()
    assert client.get(base).status_code == 401


def test_invalid_json_number_and_surrogate_return_422_not_500(knowledge):
    client, headers, project, base, case = setup(knowledge)
    for value in ["NaN", "Infinity", '"\\ud800"']:
        content = (
            '{"request_id":"' + str(uuid4()) + '","expected_revision":1,'
            '"values":{"name":"Synthetic","amount":' + value + "}}"
        )
        response = client.post(
            base + "/" + case["id"] + "/runs",
            content=content,
            headers={**headers, "Content-Type": "application/json"},
        )
        assert response.status_code == 422
    assert client.get(base + "/" + case["id"] + "/runs").json() == []


def test_standalone_creation_without_projects_and_global_exports(knowledge):
    app, client, headers, _, _ = knowledge
    base = "/api/use-cases"
    created = legacy(app, spec())
    assert created.status_code == 200, created.text
    case = created.json()
    assert case["project_id"] is None
    assert client.get(base).json() == [case]
    assert client.get(base + "/" + case["id"]).json() == case
    result = submit(client, headers, base, case, amount=12.5)
    assert result.status_code == 200, result.text
    assert client.get(base + "/" + case["id"] + "/runs").json()[0]["id"] == result.json()["id"]
    assert client.get(base + "/" + case["id"] + "/runs/" + result.json()["id"]).status_code == 200
    assert "12.5" in client.get(base + "/" + case["id"] + "/output.csv").text
    workbook = load_workbook(
        io.BytesIO(client.get(base + "/" + case["id"] + "/output.xlsx").content)
    )
    assert workbook.active["A2"].value == 12.5
    assert client.post(base, json={"spec": spec()}).status_code == 403
    client.cookies.clear()
    assert client.get(base).status_code == 401


def test_binding_changes_cas_legacy_scope_and_preserve_runs(knowledge):
    client, headers, project, legacy, case = setup(knowledge)
    base = "/api/use-cases"
    url = base + "/" + case["id"]
    original = submit(client, headers, legacy, case).json()
    detached = client.put(
        url,
        json={
            "spec": {**spec(), "execution_kind": "ai", "processing_prompt": "Synthetic processing"},
            "project_id": None,
            "expected_revision": 1,
        },
        headers=headers,
    )
    assert detached.status_code == 200, detached.text
    assert detached.json()["project_id"] is None
    assert detached.json()["revision"] == 2
    assert client.get(legacy).json() == []
    assert client.get(legacy + "/" + case["id"] + "/runs").status_code == 404
    assert client.get(url + "/runs").json()[0]["id"] == original["id"]
    assert (
        client.put(
            url,
            json={"spec": spec(), "project_id": project["id"], "expected_revision": 1},
            headers=headers,
        ).status_code
        == 409
    )
    rebound = client.put(
        url,
        json={
            "spec": {**spec(), "execution_kind": "ai", "processing_prompt": "Synthetic processing"},
            "project_id": project["id"],
            "expected_revision": 2,
        },
        headers=headers,
    )
    assert rebound.status_code == 200, rebound.text
    assert rebound.json()["project_id"] == project["id"]
    unchanged = client.put(
        url,
        json={
            "spec": {**spec(), "execution_kind": "ai", "processing_prompt": "Synthetic processing"},
            "expected_revision": 3,
        },
        headers=headers,
    )
    assert unchanged.json()["project_id"] == project["id"]
    assert (
        client.put(
            legacy + "/" + case["id"],
            json={"spec": spec(), "project_id": None, "expected_revision": 4},
            headers=headers,
        ).status_code
        == 422
    )
    assert client.get(legacy + "/" + case["id"] + "/runs").json()[0]["id"] == original["id"]


def test_optional_binding_rejects_invalid_deleted_and_archived_projects(knowledge):
    client, headers, project, _, case = setup(knowledge)
    url = "/api/use-cases/" + case["id"]
    assert (
        client.post(
            "/api/use-cases", json={"spec": spec(), "project_id": "bad"}, headers=headers
        ).status_code
        == 422
    )
    assert (
        client.put(
            url,
            json={
                "spec": {
                    **spec(),
                    "execution_kind": "ai",
                    "processing_prompt": "Synthetic processing",
                },
                "project_id": str(uuid4()),
                "expected_revision": 1,
            },
            headers=headers,
        ).status_code
        == 404
    )
    client.patch(
        "/api/projects/" + project["id"],
        json={"name": "Synthetic forms", "revision": 1, "archived": True},
        headers=headers,
    )
    assert (
        client.post(
            "/api/use-cases", json={"spec": spec(), "project_id": project["id"]}, headers=headers
        ).status_code
        == 409
    )
    assert (
        client.put(
            url,
            json={
                "spec": {
                    **spec(),
                    "execution_kind": "ai",
                    "processing_prompt": "Synthetic processing",
                },
                "project_id": None,
                "expected_revision": 1,
            },
            headers=headers,
        ).status_code
        == 409
    )
    assert submit(client, headers, "/api/use-cases", case).status_code == 409
    assert client.get(url).status_code == 200


@pytest.mark.parametrize(
    "suffix",
    ["", "/{id}", "/{id}/runs", "/{id}/runs/{run_id}", "/{id}/output.csv", "/{id}/output.xlsx"],
)
def test_read_snapshot_survives_delete_after_sql_fetch(knowledge, suffix):
    """Pause after DB materialization, delete concurrently, then decode the read response."""
    from concurrent.futures import ThreadPoolExecutor
    from threading import Event

    from sqlalchemy import event

    app, client, headers, _, _ = knowledge
    case = legacy(app, spec()).json()
    run = submit(client, headers, "/api/use-cases", case, amount=42).json()
    reached = Event()
    resume = Event()

    def pause_after_snapshot(state):
        result = state.invoke_statement()
        if state.is_select and "use_cases" in str(state.statement) and not reached.is_set():
            # Materialize and close the cursor before concurrent DELETE, preserving its snapshot.
            frozen = result.freeze()
            reached.set()
            assert resume.wait(10), "Concurrent deletion did not release the paused reader"
            return frozen()
        return result

    session_class = app.state.sessions.class_
    event.listen(session_class, "do_orm_execute", pause_after_snapshot, retval=True)
    path = "/api/use-cases" + suffix.format(id=case["id"], run_id=run["id"])
    try:
        with ThreadPoolExecutor(max_workers=1) as executor:
            pending = executor.submit(client.get, path)
            assert reached.wait(10), "Reader did not fetch its Use Case snapshot"
            try:
                deleted = client.request(
                    "DELETE",
                    "/api/use-cases/" + case["id"],
                    json={"expected_revision": 1},
                    headers=headers,
                )
                assert deleted.status_code == 200, deleted.text
            finally:
                resume.set()
            response = pending.result(timeout=10)
        assert response.status_code == 200, response.text
        if suffix == "":
            assert response.json()[0]["spec"]["name"] == spec()["name"]
        elif suffix == "/{id}":
            assert response.json()["spec"]["name"] == spec()["name"]
        elif suffix.endswith("/runs"):
            assert response.json()[0]["output"]["row"] == [42, "Synthetic"]
        elif "/runs/" in suffix:
            assert response.json()["output"]["row"] == [42, "Synthetic"]
        elif suffix.endswith(".csv"):
            assert "42;Synthetic" in response.text
        else:
            assert load_workbook(io.BytesIO(response.content)).active["A2"].value == 42
        assert client.get("/api/use-cases/" + case["id"]).status_code == 404
    finally:
        resume.set()
        event.remove(session_class, "do_orm_execute", pause_after_snapshot)
