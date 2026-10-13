"""A model's duplicate summary must be repaired once, with full cost accounting."""

import asyncio
import copy
import json
import time
from dataclasses import replace

import pytest
import test_artifact_api
from hub.artifact_drafts import DraftInput, preview, verify_token
from hub.artifact_planning import office_policy
from hub.artifact_render import _spreadsheet
from hub.artifact_specs import ArtifactSpec
from hub.db import BudgetReservation, Conversation, Message, ProviderConfig, Run
from hub.policy import Privacy, check_egress
from hub.provider_settings import seed
from hub.providers import Providers
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import select


@pytest.fixture
def setup(tmp_path, monkeypatch):
    yield from test_artifact_api.setup.__wrapped__(tmp_path, monkeypatch)


GOOD = {
    "title": "Kosten",
    "sheets": [
        {
            "name": "Kosten",
            "columns": ["Position", "Euro"],
            "rows": [["A", 10], ["B", 5]],
            "totals": {"B": "sum"},
            "number_formats": {"B": "currency"},
        }
    ],
}
BAD = copy.deepcopy(GOOD)
BAD["sheets"][0]["rows"].append(["Gesamt", {"formula": "=SUM(B2:B3)"}])
GOAL = "Erstelle eine Excel-Datei mit A: 10 Euro, B: 5 Euro und einer Gesamtsumme."


@pytest.mark.parametrize("cloud", [False, True])
@pytest.mark.parametrize("repair", [False, True])
def test_duplicate_model_total_has_one_bounded_repair_and_billed_attempts(
    setup, monkeypatch, tmp_path, cloud, repair
):
    app, _, _, cfg, pid, _ = setup
    cfg = replace(cfg, cloud_enabled=cloud, openrouter_model="synthetic/cloud")
    cid = rid = row = None
    if cloud:
        with app.state.sessions() as db:
            chat = Conversation(cloud_allowed=True)
            db.add(chat)
            db.flush()
            run = Run(conversation_id=chat.id, status="routing", context_snapshot="{}")
            db.add(run)
            db.add(
                Message(
                    conversation_id=chat.id,
                    role="user",
                    content=GOAL,
                    label=Privacy.CLOUD_ALLOWED,
                )
            )
            seed(db, cfg)
            row = db.get(ProviderConfig, "openrouter-default")
            settings = json.loads(row.config)
            settings.update(
                model="synthetic/cloud",
                context_tokens=32768,
                context_chars=32000,
                output_tokens=2048,
                max_run_cost=0.05,
                monthly_budget=1,
                pricing={"input": 0.0000001, "output": 0.0000002, "verified_at": time.time()},
            )
            row.config = json.dumps(settings)
            db.commit()
            cid, rid = chat.id, run.id

    class RepairProvider(Providers):
        calls = 0

        async def stream(self, parts, roles, *, external, grants):
            self.calls += 1
            assert external is cloud
            if cloud:
                check_egress(parts, external=True, provider_id="openrouter", grants=grants)
            else:
                assert not grants and all(p.label == Privacy.LOCAL_ONLY for p in parts)
            if self.calls == 2:
                assert "SCHEMA CORRECTION" in parts[0].content
                assert "Doppelte Aggregation" in parts[0].content
                assert parts[-1].source_id == "artifact-invalid-plan"
                assert json.dumps(BAD) in parts[-1].content
            self.measured_cost, self.cost_basis = 0.001, "provider_reported"
            yield json.dumps(GOOD if repair and self.calls == 2 else BAD)

    class LocalRepairProvider:
        # Fixture protocol avoids resolving a real configured local endpoint.
        calls = 0
        stream = RepairProvider.stream

    provider = RepairProvider(cfg) if cloud else LocalRepairProvider()
    if cloud:
        monkeypatch.setattr(
            "hub.artifact_drafts.runtime_provider",
            lambda *a, **kw: (cfg, provider, {"openrouter": row}),
        )

        async def prices(*args):
            return row

        monkeypatch.setattr("hub.artifact_drafts.fresh_cloud_prices", prices)

    def generate():
        return asyncio.run(
            preview(
                app.state.sessions,
                cfg,
                None if cloud else pid,
                DraftInput(goal=GOAL, kind="xlsx"),
                provider,
                conversation_id=cid,
                cloud_run_id=rid,
                provider_kind="openrouter" if cloud else "local",
            )
        )

    if repair:
        result = generate()
        spec = ArtifactSpec.model_validate(result["spec"])
        assert spec.plan.sheets[0].rows == GOOD["sheets"][0]["rows"]
        with app.state.sessions() as db:
            receipt = json.loads(
                verify_token(
                    db,
                    cfg,
                    result["draft_token"],
                    None if cloud else pid,
                    spec,
                    conversation_id=cid,
                )
            )
            assert receipt["draft_attempts"] == 2
        monkeypatch.setattr("hub.artifact_render.cache_formula_results", lambda path: None)
        path = tmp_path / "corrected.xlsx"
        _spreadsheet(spec.plan, path)
        sheet = load_workbook(path)["Kosten"]
        assert sheet.max_row == 4
        assert sheet["B4"].value == "=SUBTOTAL(109,B2:B3)"
        assert sum(sheet[f"B{number}"].value for number in (2, 3)) == 15
    else:
        with pytest.raises(ValidationError, match="Doppelte Aggregation"):
            generate()
    assert provider.calls == 2
    if cloud:
        with app.state.sessions() as db:
            reservations = list(db.scalars(select(BudgetReservation)))
            assert len(reservations) == 2
            assert sum(r.measured for r in reservations) == pytest.approx(0.002)
            assert db.get(Run, rid).cost == pytest.approx(0.002)


def test_spreadsheet_policy_avoids_aggregate_example_and_implicit_currency():
    policy = office_policy("xlsx", "Count items", 1)
    example = json.JSONDecoder().raw_decode(policy.split("Shape: ", 1)[1])[0]
    assert example["sheets"][0]["rows"][1][2] == {"formula": "=B3*2"}
    assert example["sheets"][0]["number_formats"] == {"B": "integer", "C": "integer"}
    assert "never both" in policy
    assert "currency only for money" in policy
