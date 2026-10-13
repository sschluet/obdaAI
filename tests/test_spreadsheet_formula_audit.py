"""Regression cases for model-authored totals and downloadable formula caches."""

import shutil
import subprocess
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from hub.artifact_render import _spreadsheet
from hub.artifact_specs import SpreadsheetPlan
from hub.spreadsheet_features import cache_formula_results
from pydantic import ValidationError


def workbook(rows, *, totals=None):
    sheet = {"name": "Input", "columns": ["Item", "Amount", "Result"], "rows": rows}
    if totals is not None:
        sheet["totals"] = totals
    return SpreadsheetPlan.model_validate(
        {"kind": "xlsx", "title": "Formula audit", "sheets": [sheet]}
    )


@pytest.mark.parametrize("function", ["SUM", "AVERAGE", "COUNT", "MIN", "MAX"])
def test_automatic_total_rejects_aggregate_of_the_same_data(function):
    # The automatic result row would include both source amounts and their aggregate.
    with pytest.raises(ValidationError):
        workbook(
            [
                ["First", 12, None],
                ["Second", 24, None],
                ["Result", {"formula": f"={function}(B2:B3)"}, None],
            ],
            totals={"B": "sum"},
        )


def test_automatic_total_allows_independent_row_calculations():
    value = workbook(
        [["First", {"formula": "=C2*2"}, 12], ["Second", {"formula": "=C3*2"}, 24]],
        totals={"B": "sum"},
    )
    assert len(value.sheets[0].rows) == 2


def test_automatic_total_allows_subtotal_excel_excludes():
    value = workbook(
        [
            ["First", 12, None],
            ["Second", 24, None],
            ["Subtotal", {"formula": "=SUBTOTAL(109,B2:B3)"}, None],
        ],
        totals={"B": "sum"},
    )
    assert value.sheets[0].rows[2][1].formula == "=SUBTOTAL(109,B2:B3)"


@pytest.mark.parametrize("formula", ["=1+", "=IF(1)", "=ROUND(1,2,3)", "=SUM(1,)", "=SUM(,1)"])
def test_incomplete_expression_or_wrong_function_arity_rejected(formula):
    with pytest.raises(ValidationError):
        workbook([["First", 12, {"formula": formula}]])


def test_reference_bounds_are_checked_on_named_cross_sheet_ranges():
    with pytest.raises(ValidationError, match="außerhalb"):
        SpreadsheetPlan.model_validate(
            {
                "kind": "xlsx",
                "title": "Bounds",
                "sheets": [
                    {"name": "Source", "columns": ["Amount"], "rows": [[12], [24]]},
                    {
                        "name": "Report",
                        "columns": ["Result"],
                        "rows": [[{"formula": "=SUM(Source!A2:A4)"}]],
                    },
                ],
            }
        )


@pytest.mark.parametrize("cached_error", ["#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A"])
def test_formula_calculation_errors_do_not_replace_downloadable_workbook(
    tmp_path, monkeypatch, cached_error
):
    monkeypatch.setattr("hub.artifact_render.cache_formula_results", lambda path: None)
    path = tmp_path / "audit.xlsx"
    _spreadsheet(workbook([["First", 12, {"formula": "=B2/0"}]]), path)
    before = path.read_bytes()
    tag = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"

    def calculate(command, **kwargs):
        target = Path(command[command.index("--outdir") + 1]) / path.name
        with zipfile.ZipFile(path) as source, zipfile.ZipFile(target, "w") as output:
            for name in source.namelist():
                data = source.read(name)
                if name == "xl/worksheets/sheet1.xml":
                    root = ET.fromstring(data)
                    for cell in root.iter(tag + "c"):
                        if cell.attrib.get("r") == "C2":
                            cell.attrib["t"] = "e"
                            cell.find(tag + "v").text = cached_error
                    data = ET.tostring(root)
                output.writestr(name, data)
        return subprocess.CompletedProcess(command, 0)

    monkeypatch.setattr("subprocess.run", calculate)
    with pytest.raises(ValueError, match="nicht berechenbar: Input!C2"):
        cache_formula_results(path)
    assert path.read_bytes() == before


@pytest.mark.skipif(not shutil.which("soffice"), reason="Offline Calc not installed locally")
def test_offline_calc_totals_ignore_subtotal_but_keep_detail_formulas(tmp_path):
    from openpyxl import load_workbook

    path = tmp_path / "totals.xlsx"
    value = workbook(
        [
            ["First", {"formula": "=C2*2"}, 12],
            ["Second", {"formula": "=C3*2"}, 24],
            ["Subtotal", {"formula": "=SUBTOTAL(109,B2:B3)"}, None],
        ],
        totals={"B": "sum"},
    )
    _spreadsheet(value, path)
    calculated = load_workbook(path, data_only=True)
    assert calculated["Input"]["B2"].value == 24
    assert calculated["Input"]["B3"].value == 48
    assert calculated["Input"]["B4"].value == 72
    assert calculated["Input"]["B5"].value == 72
    calculated.close()


@pytest.mark.skipif(not shutil.which("soffice"), reason="Offline Calc not installed locally")
def test_offline_calc_rounding_conditionals_and_negative_values(tmp_path):
    from openpyxl import load_workbook

    value = workbook(
        [
            ["Income", 120, {"formula": "=ROUND(2.675,2)"}],
            ["Expense", -24, {"formula": "=ROUND(-1.5,0)"}],
            ["Income", 80, {"formula": '=SUMIF(A2:A4,"Income",B2:B4)'}],
            ["Recovered division", None, {"formula": "=IFERROR(B2/0,0)"}],
        ],
        totals={"B": "sum"},
    )
    path = tmp_path / "financial.xlsx"
    _spreadsheet(value, path)
    calculated = load_workbook(path, data_only=True)
    sheet = calculated["Input"]
    assert sheet["C2"].value == pytest.approx(2.68)
    assert sheet["C3"].value == -2
    assert sheet["C4"].value == 200
    assert sheet["C5"].value == 0
    assert sheet["B6"].value == 176
    calculated.close()


@pytest.mark.parametrize(
    "formula",
    [
        '=SUMIF(A2:A3,"First",B2:B2)',
        '=AVERAGEIF(A2:A3,"First",B2:B2)',
        '=SUMIFS(B2:B3,A2:A2,"First")',
        '=AVERAGEIFS(B2:B3,A2:A2,"First")',
        '=COUNTIFS(A2:A3,"First",B2:B2,12)',
        "=SUMPRODUCT(B2:B3,B2:B2)",
        '=VLOOKUP("First",A2:B3,3,FALSE)',
        '=VLOOKUP("First",A2:B3,1.5,FALSE)',
        '=HLOOKUP("First",A2:B3,3,FALSE)',
        "=SUBTOTAL(112,B2:B3)",
        "=SUBTOTAL(1.5,B2:B3)",
        "=1e999",
    ],
)
def test_function_ranges_indexes_and_numeric_literals_are_checked(formula):
    with pytest.raises(ValidationError):
        workbook([["First", 12, {"formula": formula}], ["Second", 24, None]])


@pytest.mark.parametrize("formula", ["=SUM(B2:B3)*0.2", "=ROUND(SUM(B2:B3)*0.2,2)", "=SUM(B2*0.2)"])
def test_derived_tax_line_is_not_mistaken_for_duplicate_summary(formula):
    from hub.artifact_specs import SpreadsheetPlan

    SpreadsheetPlan.model_validate(
        {
            "kind": "xlsx",
            "title": "Invoice",
            "sheets": [
                {
                    "name": "Invoice",
                    "columns": ["Item", "Amount"],
                    "rows": [["A", 10], ["B", 5], ["Tax", {"formula": formula}]],
                    "totals": {"B": "sum"},
                }
            ],
        }
    )


def test_rounded_summary_is_still_rejected_with_automatic_total():
    from hub.artifact_specs import SpreadsheetPlan
    from pydantic import ValidationError

    with pytest.raises(ValidationError, match="Doppelte Aggregation"):
        SpreadsheetPlan.model_validate(
            {
                "kind": "xlsx",
                "title": "Invoice",
                "sheets": [
                    {
                        "name": "Invoice",
                        "columns": ["Item", "Amount"],
                        "rows": [
                            ["A", 10],
                            ["B", 5],
                            ["Summary", {"formula": "=ROUND(SUM(B2:B3),2)"}],
                        ],
                        "totals": {"B": "sum"},
                    }
                ],
            }
        )
