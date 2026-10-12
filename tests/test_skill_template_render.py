"""Synthetic masters only; user skill packages and originals stay outside Git."""

import hashlib
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from hub.artifact_render import render
from hub.artifact_specs import ArtifactSpec
from hub.template_render import office_stream
from openpyxl import load_workbook
from pptx import Presentation
from pydantic import ValidationError


def binding(path: Path, *, modern: bool = False) -> dict:
    return {
        "skill_id": "00000000-0000-4000-8000-000000000001",
        "revision": 2,
        "bundle_sha256": "a" * 64,
        "asset_path": "assets/" + path.name,
        "asset_sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "style": "modern" if modern else "base",
    }


def as_template(original: Path, target: Path) -> Path:
    with zipfile.ZipFile(original) as source, zipfile.ZipFile(target, "w") as output:
        for name in source.namelist():
            data = source.read(name)
            if name == "[Content_Types].xml":
                data = data.replace(
                    b"wordprocessingml.document.main", b"wordprocessingml.template.main"
                )
                data = data.replace(
                    b"presentationml.presentation.main", b"presentationml.template.main"
                )
            output.writestr(name, data)
    return target


def document_plan() -> dict:
    return {
        "kind": "docx",
        "title": "Demo strategy",
        "sections": [
            {
                "heading": "Goals",
                "paragraphs": ["Synthetic content"],
                "table": [["A", "B"], ["1", "2"]],
            }
        ],
    }


def test_dotx_preserves_headers_footers_styles_and_removes_samples(tmp_path):
    original = tmp_path / "original.docx"
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Synthetic company"
    document.sections[0].footer.paragraphs[0].text = "Synthetic footer"
    document.add_paragraph("SAMPLE BODY")
    document.styles.add_style("Synthetic Corporate Table", WD_STYLE_TYPE.TABLE)
    sample_table = document.add_table(rows=1, cols=2)
    sample_table.style = "Synthetic Corporate Table"
    document.styles["Heading 1"].font.name = "Liberation Sans"
    document.save(original)
    template = as_template(original, tmp_path / "template.dotx")
    before = template.read_bytes()
    result = render(
        ArtifactSpec(plan=document_plan(), template=binding(template)),
        [template],
        tmp_path / "output",
    )
    check = Document(tmp_path / "output" / result["filename"])
    assert check.sections[0].header.paragraphs[0].text == "Synthetic company"
    assert check.sections[0].footer.paragraphs[0].text == "Synthetic footer"
    assert check.styles["Heading 1"].font.name == "Liberation Sans"
    assert check.tables[0].style.name == "Synthetic Corporate Table"
    assert "SAMPLE BODY" not in "\n".join(p.text for p in check.paragraphs)
    assert "Goals" in "\n".join(p.text for p in check.paragraphs)
    assert template.read_bytes() == before


def test_cover_and_back_section_are_preserved(tmp_path):
    document = Document()
    cover = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    paragraph = OxmlElement("w:p")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "[Document Template]"
    run.append(text)
    paragraph.append(run)
    content.append(paragraph)
    cover.append(content)
    document.element.body.insert(0, cover)
    document.add_paragraph("Formatting sample")
    document.add_section()
    document.add_paragraph("Synthetic back cover")
    original = tmp_path / "original.docx"
    document.save(original)
    template = as_template(original, tmp_path / "cover.dotx")
    render(
        ArtifactSpec(plan=document_plan(), template=binding(template)), [template], tmp_path / "out"
    )
    check = Document(tmp_path / "out/artifact.docx")
    all_text = "".join(check.element.body.itertext())
    assert "Demo strategy" in all_text
    assert "Formatting sample" not in all_text
    assert "Synthetic back cover" in all_text
    assert len(check.sections) == 2


@pytest.mark.parametrize("modern", [False, True])
def test_potx_keeps_master_and_layout_uses_requested_style(tmp_path, modern):
    original = tmp_path / "original.pptx"
    deck = Presentation()
    deck.slide_width = 12192000
    deck.slides.add_slide(deck.slide_layouts[1]).shapes.title.text = "SAMPLE SLIDE"
    deck.save(original)
    template = as_template(original, tmp_path / "template.potx")
    spec = ArtifactSpec(
        plan={
            "kind": "pptx",
            "title": "Synthetic deck",
            "slides": [
                {"title": "New title", "bullets": ["Goal one", "Goal two"], "notes": "Speaker note"}
            ],
        },
        template=binding(template, modern=modern),
    )
    render(spec, [template], tmp_path / "out")
    check = Presentation(tmp_path / "out/artifact.pptx")
    assert check.slide_width == 12192000
    assert len(check.slides) == 1
    assert check.slides[0].shapes.title.text == "New title"
    assert check.slides[0].notes_slide.notes_text_frame.text == "Speaker note"
    assert "SAMPLE SLIDE" not in "\n".join(
        s.text for slide in check.slides for s in slide.shapes if s.has_text_frame
    )
    assert len(check.slide_layouts) == len(deck.slide_layouts)
    with (
        zipfile.ZipFile(template) as before,
        zipfile.ZipFile(tmp_path / "out/artifact.pptx") as after,
    ):
        assert before.read("ppt/slideMasters/slideMaster1.xml") == after.read(
            "ppt/slideMasters/slideMaster1.xml"
        )


def test_verified_template_required_and_old_canonical_spec_unchanged(tmp_path):
    spec = ArtifactSpec(plan=document_plan())
    assert "template" not in spec.model_dump()
    assert "letter" not in spec.model_dump()["plan"]
    template = tmp_path / "template.dotx"
    Document().save(template)
    stale = binding(template)
    stale["asset_sha256"] = "0" * 64
    with pytest.raises(ValueError, match="Vorlage wurde verändert"):
        render(ArtifactSpec(plan=document_plan(), template=stale), [template], tmp_path / "out")
    with pytest.raises(ValueError, match="Quellenbindung"):
        render(ArtifactSpec(plan=document_plan(), template=binding(template)), [], tmp_path / "out")
    for bad in ["../private.dotx", "/private.dotx", "assets/../x.dotx", "assets\\x.dotx"]:
        invalid = binding(template)
        invalid["asset_path"] = bad
        with pytest.raises(ValidationError):
            ArtifactSpec(plan=document_plan(), template=invalid)


def test_active_office_parts_are_rejected(tmp_path):
    path = tmp_path / "template.dotx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("word/embeddings/unsafe.bin", b"unsafe")
    with pytest.raises(ValueError, match="Aktiver Office-Inhalt"):
        office_stream(path, "docx")


def test_external_office_images_are_rejected(tmp_path):
    path = tmp_path / "template.potx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr(
            "ppt/_rels/x.xml.rels",
            b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            b'<Relationship Id="rId1" Type="image" Target="https://example.invalid/x.png" '
            b'TargetMode="External"/></Relationships>',
        )
    with pytest.raises(ValueError, match="Externe Office-Inhalte"):
        office_stream(path, "pptx")


def test_native_xlsx_literals_numbers_and_bounds(tmp_path):
    spec = ArtifactSpec(
        plan={
            "kind": "xlsx",
            "title": "Demo",
            "sheets": [
                {
                    "name": "Overview",
                    "columns": ["Description", "Amount"],
                    "rows": [['=HYPERLINK("https://example.invalid")', 123.5], ["Total", 20]],
                }
            ],
        }
    )
    result = render(spec, [], tmp_path / "out")
    assert result["filename"] == "artifact.xlsx"
    book = load_workbook(tmp_path / "out/artifact.xlsx")
    assert book.active["A2"].data_type == "s"
    assert book.active["B2"].value == 123.5
    assert book.active.freeze_panes == "A2"
    for name in ["bad/name", "Bad[1]"]:
        with pytest.raises(ValidationError):
            ArtifactSpec(
                plan={
                    "kind": "xlsx",
                    "title": "Demo",
                    "sheets": [{"name": name, "columns": ["X"], "rows": [[1]]}],
                }
            )
    with pytest.raises(ValidationError):
        ArtifactSpec(
            plan={
                "kind": "xlsx",
                "title": "Demo",
                "sheets": [{"name": "X", "columns": ["X"], "rows": [[float("nan")]]}],
            }
        )


def test_title_footer_control_replaced_without_losing_page_field(tmp_path):
    document = Document()
    footer = document.sections[0].footer.paragraphs[0]._p
    control = OxmlElement("w:sdt")
    properties = OxmlElement("w:sdtPr")
    alias = OxmlElement("w:alias")
    alias.set(qn("w:val"), "Title")
    properties.append(alias)
    control.append(properties)
    content = OxmlElement("w:sdtContent")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "[Document Template]"
    run.append(text)
    content.append(run)
    control.append(content)
    footer.append(control)
    run = OxmlElement("w:r")
    instruction = OxmlElement("w:instrText")
    instruction.text = "PAGE"
    run.append(instruction)
    footer.append(run)
    original = tmp_path / "original.docx"
    document.save(original)
    template = as_template(original, tmp_path / "footer.dotx")
    render(
        ArtifactSpec(plan=document_plan(), template=binding(template)), [template], tmp_path / "out"
    )
    check = Document(tmp_path / "out/artifact.docx")
    root = check.sections[0].footer._element
    assert "[Document Template]" not in "".join(root.itertext())
    assert "Demo strategy" in "".join(root.itertext())
    assert "PAGE" in [node.text for node in root.iter(qn("w:instrText"))]


def test_letter_fills_addresses_subject_and_removes_samplebody(tmp_path):
    document = Document()
    document.add_paragraph("Sample sender")
    address = document.add_table(rows=6, cols=2)
    address.cell(0, 0).text = "Sample receiver"
    address.cell(0, 1).text = "Sample author"
    subject = document.add_table(rows=2, cols=3)
    subject.cell(0, 0).text = "[Subject]"
    document.add_paragraph("Lorem ipsum")
    original = tmp_path / "original.docx"
    document.save(original)
    template = as_template(original, tmp_path / "letter.dotx")
    plan = document_plan()
    plan["letter"] = {
        "recipient": ["Synthetic receiver", "Demo road"],
        "sender": ["Synthetic sender"],
        "subject": "Demo letter",
        "date": "12.10.2026",
        "salutation": "Dear reader,",
        "signature": ["Kind regards", "Synthetic sender"],
    }
    render(ArtifactSpec(plan=plan, template=binding(template)), [template], tmp_path / "out")
    check = Document(tmp_path / "out/artifact.docx")
    assert check.tables[0].cell(0, 0).text == "Synthetic receiver"
    assert check.tables[0].cell(0, 1).text == "Synthetic sender"
    assert check.tables[1].cell(0, 0).text == "Demo letter"
    assert check.tables[1].cell(1, 2).text == "12.10.2026"
    assert "Lorem ipsum" not in "\n".join(p.text for p in check.paragraphs)
    assert "Dear reader," in "\n".join(p.text for p in check.paragraphs)


def test_letter_boundary_cleanup_preserves_substantive_body_and_missing_role(tmp_path):
    document = Document()
    document.add_table(rows=6, cols=2)
    document.add_table(rows=2, cols=3)
    original = tmp_path / "original.docx"
    document.save(original)
    template = as_template(original, tmp_path / "letter.dotx")
    plan = {
        "kind": "docx",
        "title": "Einladung zum Workshop",
        "letter": {
            "recipient": ["Beispiel GmbH"],
            "sender": ["Beispiel Absender"],
            "subject": "Einladung zum Workshop",
            "date": "12.10.2026",
            "salutation": "Sehr geehrte Damen und Herren,",
            "signature": ["Beispiel Absender", "Unterzeichner ergänzen", "Funktion ergänzen"],
        },
        "sections": [
            {
                "heading": "Brieftext",
                "paragraphs": [
                    "Sehr geehrte Damen und Herren,",
                    "Wir laden Sie herzlich zu einem Workshop ein.",
                    "Beispiel Absender",
                    "Dieser Name gehört hier ausdrücklich zum Inhalt.",
                ],
            },
            {
                "heading": "Organisatorisches",
                "paragraphs": [
                    "Der genaue Termin ist noch offen.",
                    "Mit freundlichen Grüßen\nBeispiel Absender",
                    "Unterzeichner ergänzen",
                    "Funktion ergänzen",
                ],
            },
        ],
    }
    render(ArtifactSpec(plan=plan, template=binding(template)), [template], tmp_path / "out")
    paragraphs = [p.text for p in Document(tmp_path / "out/artifact.docx").paragraphs]
    assert "Brieftext" not in paragraphs
    assert "Organisatorisches" in paragraphs
    assert paragraphs.count("Sehr geehrte Damen und Herren,") == 1
    assert paragraphs.count("Mit freundlichen Grüßen") == 1
    assert (
        paragraphs.count("Beispiel Absender") == 2
    )  # One substantive body line plus one signature.
    assert paragraphs.count("Funktion ergänzen") == 1
    assert "Unterzeichner ergänzen" not in paragraphs
    assert "Dieser Name gehört hier ausdrücklich zum Inhalt." in paragraphs
    assert "Der genaue Termin ist noch offen." in paragraphs


def test_letter_missing_signer_placeholder_remains_explicit():
    from hub.template_render import _letter_parts

    plan = document_plan()
    plan["letter"] = {
        "recipient": ["Beispiel"],
        "subject": "Demo",
        "date": "12.10.2026",
        "salutation": "Hallo,",
        "signature": ["Unterzeichner ergänzen"],
    }
    parts, signature = _letter_parts(ArtifactSpec(plan=plan).plan)
    assert signature == ["Unterzeichner ergänzen"]
    assert parts[0][0] == "Goals"
