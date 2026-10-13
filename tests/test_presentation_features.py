"""Synthetic native diagrams only; no corporate files in fixtures."""

import zipfile

import pytest
from hub.artifact_specs import PresentationPlan, SlideVisual
from hub.presentation_features import content, foreground
from hub.template_render import presentation
from pptx import Presentation
from pptx.util import Inches
from pydantic import ValidationError


def examples():
    items = [
        {"label": f"Stage {i}", "text": "Synthetic detail", "value": str(i + 2)} for i in range(4)
    ]
    for kind in ("process", "cycle", "kpi", "layers", "messages", "dot_matrix"):
        yield {"kind": kind, "items": items}
    yield {
        "kind": "swimlane",
        "items": [dict(i, lane=f"Lane {n // 2}") for n, i in enumerate(items)],
    }
    yield {
        "kind": "comparison",
        "items": [dict(i, lane=f"Side {n // 2}") for n, i in enumerate(items)],
    }
    yield {
        "kind": "timeline",
        "columns": ["Q1", "Q2", "Q3"],
        "items": [dict(i, lane=f"Lane {n}", start=0, end=2.5) for n, i in enumerate(items)],
    }
    yield {
        "kind": "matrix",
        "columns": ["Q1", "Q2", "Q3"],
        "items": [dict(i, lane=f"Lane {n}", period=n % 3) for n, i in enumerate(items)],
    }
    for kind in ("chart", "status_bars"):
        yield {
            "kind": kind,
            "columns": ["A", "B", "C"],
            "series": [{"name": "Demo", "values": [2, 5, 7]}],
        }


@pytest.mark.parametrize("data", list(examples()))
def test_native_visuals_editable_and_within_grid(tmp_path, data):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    plan = PresentationPlan(
        kind="pptx",
        title="Demo",
        slides=[
            {
                "title": "Synthetic statement",
                "bullets": ["Demo"],
                "notes": "Speaker evidence",
                "visual": data,
            }
        ],
    )
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    content(slide, plan.slides[0], 1, plan.title, 1)
    assert len(slide.shapes) > 4
    for obj in slide.shapes:
        assert obj.left >= 0 and obj.top >= 0
        assert obj.left + obj.width <= deck.slide_width
        assert obj.top + obj.height <= deck.slide_height
    path = tmp_path / "demo.pptx"
    deck.save(path)
    check = Presentation(path)
    assert check.slides[0].notes_slide.notes_text_frame.text == "Speaker evidence"
    with zipfile.ZipFile(path) as archive:
        assert not any("media/image" in p for p in archive.namelist())
        assert bool(
            [
                p
                for p in archive.namelist()
                if p.startswith("ppt/charts/chart") and p.endswith(".xml")
            ]
        ) == (data["kind"] == "chart")
        xml = archive.read("ppt/slides/slide1.xml")
        assert b"gradFill" not in xml and b"outerShdw" not in xml


def test_modern_uses_named_master_layouts_and_exact_grid(tmp_path):
    template = tmp_path / "template.pptx"
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    for number, name in ((0, "Cover 1"), (6, "Empty 1"), (2, "End 2")):
        deck.slide_layouts[number].name = name
    deck.save(template)
    target = tmp_path / "result.pptx"
    plan = PresentationPlan(
        kind="pptx",
        title="Demo",
        slides=[
            {"title": "Cover", "bullets": ["Introduction"], "layout": "cover"},
            {
                "title": "Results",
                "bullets": ["Synthetic"],
                "visual": {"kind": "kpi", "items": [{"label": "Demo", "value": "25"}]},
            },
            {"title": "End", "bullets": ["Contact"], "layout": "end"},
        ],
    )
    presentation(plan, target, template, True)
    check = Presentation(target)
    assert [s.slide_layout.name for s in check.slides] == ["Cover 1", "Empty 1", "End 2"]
    assert any(s.left == Inches(0.52) and s.top == Inches(0.74) for s in check.slides[1].shapes)


def test_contrast_and_invalid_geometry():
    assert foreground("020623") == "FFFFFF"
    assert foreground("EB601D") == "020623"
    assert foreground("4778E2") == "020623"
    with pytest.raises(ValidationError):
        SlideVisual(kind="chart", columns=["A", "B"], series=[{"name": "X", "values": [1]}])
    with pytest.raises(ValidationError):
        SlideVisual(
            kind="timeline",
            columns=["A"],
            items=[{"label": "X", "lane": "Demo", "start": 0, "end": 2}],
        )


def test_short_timeline_span_becomes_editable_milestone():
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    plan = PresentationPlan(
        kind="pptx",
        title="Demo",
        slides=[
            {
                "title": "Tiny phase",
                "bullets": ["Demo"],
                "visual": {
                    "kind": "timeline",
                    "columns": ["Q1"],
                    "items": [{"label": "Close", "lane": "Demo", "start": 0.999, "end": 1}],
                },
            }
        ],
    )
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    content(slide, plan.slides[0], 1, "Demo", 1)
    assert all(s.width >= 0 and s.left + s.width <= deck.slide_width for s in slide.shapes)
    assert any(s.has_text_frame and s.text == "Close" for s in slide.shapes)


def test_max_dense_text_is_rejected_instead_of_overlapping():
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    plan = PresentationPlan(
        kind="pptx",
        title="Demo",
        slides=[
            {
                "title": "Demo",
                "bullets": ["Demo"],
                "visual": {
                    "kind": "swimlane",
                    "items": [
                        {
                            "label": "Long word " * 8,
                            "text": "Dense content " * 15,
                            "lane": str(i // 6),
                        }
                        for i in range(24)
                    ],
                },
            }
        ],
    )
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    with pytest.raises(ValueError, match="Text kürzen"):
        content(slide, plan.slides[0], 1, "Demo", 1)


def test_end_summary_is_not_duplicated(tmp_path):
    deck = Presentation()
    deck.slide_layouts[2].name = "End 2"
    template = tmp_path / "master.pptx"
    deck.save(template)
    target = tmp_path / "end.pptx"
    plan = PresentationPlan(
        kind="pptx",
        title="Demo",
        slides=[{"title": "End", "layout": "end", "bullets": ["Unique conclusion"]}],
    )
    presentation(plan, target, template, True)
    check = Presentation(target)
    texts = [s.text for s in check.slides[0].shapes if s.has_text_frame]
    assert sum(t.count("Unique conclusion") for t in texts) == 1


def test_semantic_periods_resolve_and_overlaps_pack_without_mutation():
    from hub.presentation_features import period_bounds, timeline_rows

    plan = SlideVisual(
        kind="timeline",
        columns=["Q1", "Q2", "Q3"],
        items=[
            {"label": "Task A", "lane": "Delivery", "from_period": "Q1", "to_period": "Q2"},
            {"label": "Task B", "lane": "Delivery", "from_period": "Q2", "to_period": "Q3"},
            {"label": "Task C", "lane": "Delivery", "from_period": "Q3", "to_period": "Q3"},
        ],
    )
    original = plan.model_dump_json()
    assert period_bounds(plan, plan.items[0]) == (0, 2)
    assert period_bounds(plan, plan.items[2]) == (2, 3)
    rows = timeline_rows(plan)
    assert [[i.label for i in row] for _, row in rows] == [["Task A", "Task C"], ["Task B"]]
    assert plan.model_dump_json() == original
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    from hub.presentation_features import visual

    visual(slide, plan, (0.52, 1.85, 12.29, 4.7))
    texts = [s.text for s in slide.shapes if s.has_text_frame]
    assert all(texts.count(label) == 1 for label in ("Task A", "Task B", "Task C"))


@pytest.mark.parametrize(
    "periods",
    [
        {"from_period": "Q1"},
        {"from_period": "Unknown", "to_period": "Q2"},
        {"from_period": "Q2", "to_period": "Q1"},
        {"from_period": "Q1", "to_period": "Q2", "start": 0, "end": 2},
    ],
)
def test_semantic_periods_fail_closed(periods):
    with pytest.raises(ValidationError):
        SlideVisual(
            kind="timeline",
            columns=["Q1", "Q2"],
            items=[{"label": "Demo", "lane": "Delivery", **periods}],
        )


def test_overlapping_matrix_cells_use_separate_rows_and_density_bound():
    from hub.presentation_features import timeline_rows

    plan = SlideVisual(
        kind="matrix",
        columns=["Q1"],
        items=[{"label": f"Task {i}", "lane": "Delivery", "period": 0} for i in range(7)],
    )
    with pytest.raises(ValueError, match="sechs"):
        timeline_rows(plan)


@pytest.mark.parametrize("layout", ["cover", "end"])
def test_bookend_long_summary_keeps_every_visible_item(tmp_path, layout):
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    deck.slide_layouts[0 if layout == "cover" else 2].name = (
        "Cover 1" if layout == "cover" else "End 2"
    )
    template = tmp_path / "master.pptx"
    deck.save(template)
    target = tmp_path / "bookend.pptx"
    bullets = [f"Conclusion {i}: " + "Synthetic explanatory information. " * 3 for i in range(5)]
    plan = PresentationPlan(
        kind="pptx",
        title="Demo",
        slides=[
            {
                "title": "Visible conclusions",
                "layout": layout,
                "bullets": bullets,
                "subtitle": "Distinct subtitle",
                "notes": "Speaker notes",
            }
        ],
    )
    presentation(plan, target, template, True)
    slide = Presentation(target).slides[0]
    visible = "\n".join(s.text for s in slide.shapes if s.has_text_frame)
    assert all(visible.count(bullet) == 1 for bullet in bullets)
    assert visible.count("Distinct subtitle") == 1
    assert slide.notes_slide.notes_text_frame.text == "Speaker notes"
