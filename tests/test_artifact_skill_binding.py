"""Local model plans keep exact reviewed package/template provenance through dispatch."""

import io
import json

import pytest
from docx import Document
from hub.artifact_specs import ArtifactSpec
from hub.db import ArtifactRevision
from test_artifact_api import setup as setup
from test_artifact_drafts import DOCUMENT, DraftProvider
from test_skill_bundles import bundle


def imported(c, h, extra=None):
    document = Document()
    document.sections[0].header.paragraphs[0].text = "Synthetic master header"
    document.add_paragraph("Template sample")
    output = io.BytesIO()
    document.save(output)
    package = bundle({"synthetic/assets/master.docx": output.getvalue(), **(extra or {})})
    response = c.post(
        "/api/skills/import-package", headers={**h, "X-File-Name": "test.skill"}, content=package
    )
    assert response.status_code == 200, response.text
    skill = response.json()
    response = c.post(
        "/api/skills/" + skill["id"] + "/trust",
        headers=h,
        json={"expected_revision": skill["revision"], "sha256": skill["sha256"], "enabled": True},
    )
    assert response.status_code == 200
    return skill


def propose(app, c, h, pid, skill):
    app.state.providers = DraftProvider(DOCUMENT)
    result = c.post(
        "/api/projects/" + pid + "/artifact-drafts",
        headers=h,
        json={
            "goal": "Create a synthetic Word report",
            "kind": "docx",
            "skill_bindings": [
                {"id": skill["id"], "revision": skill["revision"], "sha256": skill["sha256"]}
            ],
        },
    )
    assert result.status_code == 200, result.text
    return result.json()


def test_template_enters_exact_spool_and_receipt(setup):
    app, c, h, cfg, pid, _ = setup
    skill = imported(c, h)
    proposal = propose(app, c, h, pid, skill)
    template = proposal["spec"]["template"]
    assert template["bundle_sha256"] == skill["bundle_sha256"]
    import uuid

    saved = c.post(
        "/api/projects/" + pid + "/artifacts",
        headers=h,
        json={
            "request_id": str(uuid.uuid4()),
            "spec": proposal["spec"],
            "draft_token": proposal["draft_token"],
        },
    )
    assert saved.status_code == 200, saved.text
    row = saved.json()
    with app.state.sessions() as db:
        revision = db.get(ArtifactRevision, (row["artifact_id"], row["revision"]))
        assert json.loads(revision.draft_metadata)["skills"][0]["id"] == skill["id"]
    from pathlib import Path

    from hub.artifact_render import render

    spool = Path(cfg.artifact_request_dir) / (row["artifact_id"] + "-1")
    spec = ArtifactSpec.model_validate_json((spool / "spec.json").read_text())
    manifest = render(spec, [spool / "source-0"], Path(cfg.artifact_result_dir) / "direct")
    output = Document(Path(cfg.artifact_result_dir) / "direct" / manifest["filename"])
    assert output.sections[0].header.paragraphs[0].text == "Synthetic master header"
    assert "Draft text." in "\n".join(p.text for p in output.paragraphs)


@pytest.mark.parametrize("mutation", ["template", "disable", "foreign"])
def test_template_mutation_or_revocation_rejected(setup, mutation):
    app, c, h, cfg, pid, _ = setup
    skill = imported(c, h)
    proposal = propose(app, c, h, pid, skill)
    if mutation == "template":
        proposal["spec"]["template"]["asset_sha256"] = "0" * 64
    elif mutation == "disable":
        c.post(
            "/api/skills/" + skill["id"] + "/trust",
            headers=h,
            json={"expected_revision": 1, "sha256": skill["sha256"], "enabled": False},
        )
    else:
        from hub.db import Skill

        with app.state.sessions() as db:
            db.get(Skill, skill["id"]).project_id = c.post(
                "/api/projects", headers=h, json={"name": "Other"}
            ).json()["id"]
            db.commit()
    import uuid

    result = c.post(
        "/api/projects/" + pid + "/artifacts",
        headers=h,
        json={
            "request_id": str(uuid.uuid4()),
            "spec": proposal["spec"],
            "draft_token": proposal["draft_token"],
        },
    )
    assert result.status_code == 409


def test_reference_guidance_is_hash_bound_bounded_and_not_topic_evidence(setup):
    app, c, h, cfg, pid, _ = setup
    reference = (
        "# Typography\nUse generous margins and native charts.\n"
        "```python\nSCRIPT_MUST_NOT_ENTER\n```\n"
    )
    skill = imported(
        c,
        h,
        {
            "synthetic/references/gestaltung.md": reference,
            "synthetic/scripts/helper.py": "SCRIPT_MUST_NOT_ENTER",
        },
    )
    proposal = propose(app, c, h, pid, skill)
    provider = app.state.providers
    guidance = [part.content for part in provider.parts if part.source_id.startswith("skill:")]
    assert len(guidance) == 1 and len(guidance[0]) < 5000
    assert "Use generous margins" in guidance[0]
    assert "NOT factual topic evidence" in guidance[0]
    assert "SCRIPT_MUST_NOT_ENTER" not in guidance[0]
    assert "NEVER turn their sample business narratives" in provider.parts[0].content
    import uuid

    response = c.post(
        "/api/projects/" + pid + "/artifacts",
        headers=h,
        json={
            "request_id": str(uuid.uuid4()),
            "spec": proposal["spec"],
            "draft_token": proposal["draft_token"],
        },
    )
    assert response.status_code == 200, response.text
    with app.state.sessions() as db:
        row = db.get(ArtifactRevision, (response.json()["artifact_id"], 1))
        references = json.loads(row.draft_metadata)["skills"][0]["guidance_files"]
        assert references == [
            {
                "path": "synthetic/references/gestaltung.md",
                "sha256": __import__("hashlib").sha256(reference.encode()).hexdigest(),
            }
        ]


def test_altered_reference_manifest_rejects_draft_receipt(setup):
    app, c, h, cfg, pid, _ = setup
    skill = imported(c, h, {"synthetic/references/layouts.md": "# Grid\nNative chart layout."})
    proposal = propose(app, c, h, pid, skill)
    from hub.db import SkillRevision

    with app.state.sessions() as db:
        revision = db.get(SkillRevision, (skill["id"], skill["revision"]))
        manifest = json.loads(revision.bundle_manifest)
        for resource in manifest["files"]:
            if resource["path"].endswith("layouts.md"):
                resource["sha256"] = "0" * 64
        revision.bundle_manifest = json.dumps(manifest)
        db.commit()
    import uuid

    response = c.post(
        "/api/projects/" + pid + "/artifacts",
        headers=h,
        json={
            "request_id": str(uuid.uuid4()),
            "spec": proposal["spec"],
            "draft_token": proposal["draft_token"],
        },
    )
    assert response.status_code == 409


def test_sdk_prose_is_not_content_but_reviewed_custom_edits_survive(setup):
    from hub.artifact_skill_binding import context
    from hub.db import Skill, SkillRevision
    from hub.skill_api import parse, portable, sha

    app, c, h, cfg, pid, _spec = setup
    data = bundle(
        front="name: avendis-excel\ndescription: Synthetic reviewed Excel adapter",
        body="# SDK\nRASTER_SDK_SLOGAN must never become spreadsheet content.",
    )
    response = c.post(
        "/api/skills/import-package", headers={**h, "X-File-Name": "demo.skill"}, content=data
    )
    assert response.status_code == 200
    skill = response.json()
    assert (
        c.post(
            "/api/skills/" + skill["id"] + "/trust",
            headers=h,
            json={
                "expected_revision": skill["revision"],
                "sha256": skill["sha256"],
                "enabled": True,
            },
        ).status_code
        == 200
    )
    binding = {"id": skill["id"], "revision": skill["revision"], "sha256": skill["sha256"]}
    with app.state.sessions() as db:
        parts, _, _ = context(db, cfg, pid, [binding], "xlsx", "Calculate supplied numbers")
        assert "RASTER_SDK_SLOGAN" not in parts[0].content
        row = db.get(Skill, skill["id"])
        old = db.get(SkillRevision, (row.id, row.revision))
        meta, body = parse(old.document)
        document = portable(meta, body + "\n\nUse a dedicated Summary sheet for custom results.")
        row.revision += 1
        row.trusted_sha256 = sha(document)
        db.add(
            SkillRevision(
                skill_id=row.id,
                revision=row.revision,
                document=document,
                sha256=row.trusted_sha256,
                bundle_sha256=old.bundle_sha256,
                bundle_manifest=old.bundle_manifest,
            )
        )
        db.commit()
        parts, _, _ = context(
            db,
            cfg,
            pid,
            [{"id": row.id, "revision": row.revision, "sha256": row.trusted_sha256}],
            "xlsx",
            "Calculate supplied numbers",
        )
        assert "dedicated Summary sheet" in parts[0].content
        assert "RASTER_SDK_SLOGAN" not in parts[0].content
