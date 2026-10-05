import hashlib
from pathlib import Path

import pytest
from docx import Document
from hub.artifact_render import render
from hub.artifact_specs import ArtifactSpec
from PIL import Image
from pptx import Presentation
from pydantic import ValidationError
from pypdf import PdfReader, PdfWriter


def binding(path: Path) -> dict[str, object]:
    return {
        "id": "11111111-1111-4111-8111-111111111111",
        "revision": 1,
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
    }


def test_word_is_editable_and_version_is_immutable(tmp_path: Path) -> None:
    spec = ArtifactSpec.model_validate(
        {
            "plan": {
                "kind": "docx",
                "title": "Lokaler Bericht – Übersicht",
                "sections": [
                    {
                        "heading": "Ergebnis",
                        "paragraphs": ["Überprüfbarer Inhalt."],
                        "table": [["Name", "Status"], ["Test", "Erfolgreich"]],
                    }
                ],
            }
        }
    )
    out = tmp_path / "v1"
    manifest = render(spec, [], out)
    doc = Document(str(out / "artifact.docx"))
    assert doc.paragraphs[2].text == "Überprüfbarer Inhalt."
    assert doc.tables[0].cell(1, 1).text == "Erfolgreich"
    doc.paragraphs[2].text = "Editierbarer neuer Inhalt"
    doc.save(str(tmp_path / "edited.docx"))
    assert manifest["label"] == "LOCAL_ONLY" and not manifest["preview_verified"]
    assert manifest["sha256"] == hashlib.sha256((out / "artifact.docx").read_bytes()).hexdigest()
    before = (out / "artifact.docx").read_bytes()
    with pytest.raises(ValueError, match="existiert"):
        render(spec, [], out)
    assert (out / "artifact.docx").read_bytes() == before


def test_powerpoint_has_requested_count_and_editable_notes(tmp_path: Path) -> None:
    spec = ArtifactSpec.model_validate(
        {
            "plan": {
                "kind": "pptx",
                "title": "Projektstand",
                "slides": [
                    {"title": "Stand", "bullets": ["Eine echte Folie"], "notes": "Sprechernotiz"},
                    {"title": "Nächste Schritte", "bullets": ["Weiterarbeiten", "Ergebnis prüfen"]},
                ],
            }
        }
    )
    render(spec, [], tmp_path / "slides")
    deck = Presentation(str(tmp_path / "slides/artifact.pptx"))
    assert len(deck.slides) == 2
    for slide in deck.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                assert all(p.font.name == "Liberation Sans" for p in shape.text_frame.paragraphs)
    assert "Eine echte Folie" in deck.slides[0].shapes[1].text
    assert deck.slides[0].notes_slide.notes_text_frame.text == "Sprechernotiz"
    deck.slides[0].shapes[0].text = "Geändert"
    deck.save(str(tmp_path / "edited.pptx"))
    assert Presentation(str(tmp_path / "edited.pptx")).slides[0].shapes[0].text == "Geändert"


def test_pdf_reorders_rotates_removes_and_preserves_original(tmp_path: Path) -> None:
    source = tmp_path / "input.pdf"
    writer = PdfWriter()
    for width in [100, 200, 300]:
        writer.add_blank_page(width=width, height=400)
    writer.add_js("app.alert('never execute');")
    writer.write(str(source))
    before = source.read_bytes()
    spec = ArtifactSpec.model_validate(
        {
            "sources": [binding(source)],
            "plan": {
                "kind": "pdf",
                "title": "Ausgewählte Seiten",
                "pages": [{"source": 0, "page": 3, "rotation": 90}, {"source": 0, "page": 1}],
            },
        }
    )
    render(spec, [source], tmp_path / "pdf")
    result = PdfReader(str(tmp_path / "pdf/artifact.pdf"))
    assert len(result.pages) == 2
    assert result.pages[0].mediabox.width == 300 and result.pages[0].rotation == 90
    assert result.pages[1].mediabox.width == 100
    assert "/Names" not in result.trailer["/Root"]
    assert "/OpenAction" not in result.trailer["/Root"]
    assert source.read_bytes() == before


def test_crop_collage_changes_pixels_and_strips_metadata(tmp_path: Path) -> None:
    source = tmp_path / "input.png"
    original = Image.new("RGB", (256, 128), "red")
    original.paste("blue", (128, 0, 256, 128))
    original.save(source)
    spec = ArtifactSpec.model_validate(
        {
            "sources": [binding(source)],
            "plan": {
                "kind": "image",
                "title": "Zwei Ausschnitte",
                "columns": 2,
                "tile_width": 128,
                "tile_height": 128,
                "tiles": [
                    {"source": 0, "crop": [0, 0, 0.5, 1]},
                    {"source": 0, "crop": [0.5, 0, 1, 1]},
                ],
            },
        }
    )
    render(spec, [source], tmp_path / "image")
    with Image.open(tmp_path / "image/artifact.png") as output:
        assert output.size == (256, 128)
        assert output.getpixel((64, 64)) == (255, 0, 0)
        assert output.getpixel((192, 64)) == (0, 0, 255)
        assert not output.getexif()


def test_source_drift_and_symlinks_fail_before_output(tmp_path: Path) -> None:
    source = tmp_path / "source.png"
    Image.new("RGB", (128, 128)).save(source)
    spec = ArtifactSpec.model_validate(
        {
            "sources": [binding(source)],
            "plan": {
                "kind": "image",
                "title": "Bild",
                "tiles": [{"source": 0}],
            },
        }
    )
    alias = tmp_path / "linked.png"
    alias.symlink_to(source)
    with pytest.raises(ValueError, match="Verknüpfte"):
        render(spec, [alias], tmp_path / "bad-link")
    source.write_bytes(b"changed")
    with pytest.raises(ValueError, match="verändert"):
        render(spec, [source], tmp_path / "bad-hash")
    assert not (tmp_path / "bad-link").exists() and not (tmp_path / "bad-hash").exists()


def test_invalid_pdf_page_cleans_failed_output(tmp_path: Path) -> None:
    source = tmp_path / "source.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(str(source))
    spec = ArtifactSpec.model_validate(
        {
            "sources": [binding(source)],
            "plan": {
                "kind": "pdf",
                "title": "PDF",
                "pages": [{"source": 0, "page": 2}],
            },
        }
    )
    with pytest.raises(ValueError, match="existiert nicht"):
        render(spec, [source], tmp_path / "missing-page")
    assert not (tmp_path / "missing-page").exists()


@pytest.mark.parametrize("extra", ["code", "url", "path", "command", "cloud_allowed"])
def test_no_executable_or_external_plan_fields(extra: str) -> None:
    with pytest.raises(ValidationError):
        ArtifactSpec.model_validate(
            {
                "plan": {
                    "kind": "docx",
                    "title": "Test",
                    "sections": [{"heading": "Ergebnis"}],
                    extra: "external or executable",
                }
            }
        )


def test_structural_limits_and_local_only() -> None:
    with pytest.raises(ValidationError):
        ArtifactSpec.model_validate(
            {
                "label": "CLOUD_ALLOWED",
                "plan": {
                    "kind": "docx",
                    "title": "Test",
                    "sections": [{"heading": "Ergebnis"}],
                },
            }
        )
    with pytest.raises(ValidationError):
        ArtifactSpec.model_validate(
            {
                "plan": {
                    "kind": "pptx",
                    "title": "Test",
                    "slides": [{"title": "Überladen", "bullets": ["x"] * 6}],
                }
            }
        )
    with pytest.raises(ValidationError):
        ArtifactSpec.model_validate(
            {
                "plan": {
                    "kind": "docx",
                    "title": "Test",
                    "sections": [{"heading": "Tabelle", "table": [["a"], ["b", "c"]]}],
                }
            }
        )


def test_document_embeds_bound_image_with_caption_without_changing_original(tmp_path):
    source = tmp_path / "synthetic.png"
    Image.new("RGB", (180, 120), "blue").save(source)
    before = source.read_bytes()
    spec = ArtifactSpec.model_validate(
        {
            "sources": [binding(source)],
            "plan": {
                "kind": "docx",
                "title": "Bilder im Dokument",
                "pdf_export": True,
                "sections": [
                    {
                        "heading": "Ergebnis",
                        "images": [
                            {"source": 0, "caption": "Synthetisches blaues Bild", "width_inches": 3}
                        ],
                    }
                ],
            },
        }
    )
    render(spec, [source], tmp_path / "doc-image")
    doc = Document(tmp_path / "doc-image/artifact.docx")
    assert len(doc.inline_shapes) == 1
    assert doc.inline_shapes[0].width.inches == 3
    assert doc.paragraphs[-1].text == "Synthetisches blaues Bild"
    assert source.read_bytes() == before
    with pytest.raises(ValidationError):
        ArtifactSpec.model_validate({"plan": spec.plan.model_dump()})


def test_caption_reserves_space_and_rejects_overflow(tmp_path):
    source = tmp_path / "synthetic.png"
    Image.new("RGB", (128, 128), "red").save(source)
    data = {
        "sources": [binding(source)],
        "plan": {
            "kind": "image",
            "title": "Beschriftung",
            "tile_width": 256,
            "tile_height": 128,
            "tiles": [{"source": 0, "caption": "Blau und Gelb"}],
        },
    }
    render(ArtifactSpec.model_validate(data), [source], tmp_path / "caption")
    with Image.open(tmp_path / "caption/artifact.png") as image:
        assert image.getpixel((100, 40)) == (255, 0, 0)
        assert image.getpixel((255, 127)) == (244, 247, 253)
        assert any(
            image.getpixel((x, y)) == (24, 40, 66) for x in range(8, 150) for y in range(100, 120)
        )
    data["plan"]["tiles"][0]["caption"] = "x" * 140
    with pytest.raises(ValueError, match="langes Wort"):
        render(ArtifactSpec.model_validate(data), [source], tmp_path / "overflow")


def test_absent_extensions_preserve_previous_canonical_request_shape():
    import json

    from hub.artifact_runner import canonical

    old = {
        "version": 1,
        "label": "LOCAL_ONLY",
        "sources": [],
        "plan": {
            "kind": "docx",
            "title": "Unverändert",
            "sections": [{"heading": "Abschnitt", "paragraphs": ["Text"], "table": []}],
        },
    }
    assert canonical(ArtifactSpec.model_validate(old)) == json.dumps(
        old, ensure_ascii=False, sort_keys=True
    )


def test_new_pdf_uses_fixed_office_conversion_and_cleans_temporary_docx(tmp_path, monkeypatch):
    def convert(path, folder):
        doc = Document(path)
        assert doc.paragraphs[0].text == "Neues PDF"
        assert doc.paragraphs[2].text == "Text im PDF"
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=300)
        target = folder / "artifact.pdf"
        writer.write(target)
        return target

    monkeypatch.setattr("hub.artifact_preview.office_pdf", convert)
    spec = ArtifactSpec.model_validate(
        {
            "plan": {
                "kind": "pdf",
                "title": "Neues PDF",
                "sections": [{"heading": "Abschnitt", "paragraphs": ["Text im PDF"]}],
            }
        }
    )
    render(spec, [], tmp_path / "text-pdf")
    assert len(PdfReader(tmp_path / "text-pdf/artifact.pdf").pages) == 1
    assert not (tmp_path / "text-pdf/artifact.docx").exists()
    for plan in [
        {"kind": "pdf", "title": "Leer"},
        {
            "kind": "pdf",
            "title": "Gemischt",
            "sections": [{"heading": "X"}],
            "pages": [{"source": 0, "page": 1}],
        },
    ]:
        with pytest.raises(ValidationError):
            ArtifactSpec.model_validate({"plan": plan})
