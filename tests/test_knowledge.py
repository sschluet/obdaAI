import asyncio
import hashlib
import time
import uuid
import zipfile

import pytest
from alembic import command
from alembic.config import Config
from fastapi.testclient import TestClient
from hub.api import create_app
from hub.auth import password_hash
from hub.config import Settings
from hub.db import Memory, Run, Source
from hub.extract import isolated_parse, parse, safe_archive
from hub.knowledge import context
from hub.memory import digest, reconcile, vault_path
from hub.routing import Route
from hub.websearch import SearchInput, search, target
from sqlalchemy import select
from test_chat import FakeProviders


@pytest.fixture
def knowledge(tmp_path, monkeypatch):
    url = "sqlite:///" + str(tmp_path / "db.sqlite")
    monkeypatch.setenv("HUB_DATABASE_URL", url)
    command.upgrade(Config("alembic.ini"), "head")
    vault = tmp_path / "vault"
    vault.mkdir()
    cfg = Settings(
        database_url=url,
        origin="http://testserver",
        secure_cookie=False,
        password_hash=password_hash("synthetic-knowledge-password"),
        source_dir=str(tmp_path / "sources"),
        vault_dir=str(vault),
        vault_writable=True,
    )
    app = create_app(cfg)
    app.state.providers = FakeProviders(
        Route(
            capabilities=["retrieval"],
            model_class="cloud",
            reason="synthetic source question",
            answer=None,
        )
    )
    with TestClient(app) as client:
        login = client.post(
            "/api/login",
            json={"password": "synthetic-knowledge-password"},
            headers={"Origin": "http://testserver"},
        )
        headers = {"Origin": "http://testserver", "X-CSRF-Token": login.json()["csrf"]}
        yield app, client, headers, cfg, vault


def project(client, headers, name="Synthetic project"):
    r = client.post("/api/projects", json={"name": name}, headers=headers)
    assert r.status_code == 200
    return r.json()["id"]


def upload(client, headers, pid, name, content):
    r = client.post(
        "/api/projects/" + pid + "/sources",
        content=content,
        headers={**headers, "X-File-Name": name},
    )
    assert r.status_code == 200
    return r.json()["id"]


def wait(app, sid):
    for _ in range(100):
        with app.state.sessions() as db:
            row = db.get(Source, sid)
            if row.status in {"ready", "failed"}:
                return row
        time.sleep(0.025)
    pytest.fail("Extraction did not complete")


def test_immutable_sources_scoped_retrieval_and_chat(knowledge):
    app, c, h, cfg, v = knowledge
    a, b = project(c, h, "A"), project(c, h, "B")
    sa = upload(c, h, a, "a.md", b"# invoice\nsynthetic_invoice 10 EUR")
    sb = upload(c, h, b, "b.md", b"# invoice\nsynthetic_private_other_project 500 EUR")
    assert wait(app, sa).status == wait(app, sb).status == "ready"
    found = c.get("/api/projects/" + a + "/search?q=invoice").json()
    assert found and all(x["source_id"] == sa for x in found)
    assert c.get("/api/projects/" + a + "/sources/" + sb + "/original").status_code == 404
    r = c.get("/api/projects/" + a + "/sources/" + sa + "/original")
    assert r.content == b"# invoice\nsynthetic_invoice 10 EUR"
    assert hashlib.sha256(r.content).hexdigest() == found[0]["sha256"]
    cid = c.post("/api/conversations", json={"project_id": a}, headers=h).json()["id"]
    result = c.post(
        "/api/conversations/" + cid + "/messages", json={"content": "synthetic_invoice"}, headers=h
    )
    assert "succeeded" in result.text
    assert app.state.providers.stream_calls == [False]
    with app.state.sessions() as db:
        run = db.scalar(select(Run).where(Run.conversation_id == cid))
        assert "synthetic_invoice" in run.context_snapshot
        assert "synthetic_private_other_project" not in run.context_snapshot
    assert (
        c.post(
            "/api/conversations", json={"project_id": a, "cloud_allowed": True}, headers=h
        ).status_code
        == 409
    )


def test_decimal_table_audit(knowledge):
    app, c, h, cfg, v = knowledge
    pid = project(c, h)
    sid = upload(c, h, pid, "numbers.csv", b"Name;EUR\nA;1.234,50\nB;0,10\nFormula;=SUM(B2:B3)\n")
    assert wait(app, sid).status == "ready"
    result = c.get(f"/api/projects/{pid}/sources/{sid}/sum?column=1").json()
    assert result["total"] == "1234.60"
    assert len(result["values"]) == 2 and len(result["skipped"]) == 2
    assert all("CSV Zeile" in x["locator"] for x in result["values"])


@pytest.mark.parametrize("filename", ["../secret.md", "bad.exe", "bad.docm", "bad\\path.md"])
def test_upload_paths_types_and_csrf(knowledge, filename):
    app, c, h, cfg, v = knowledge
    pid = project(c, h)
    assert (
        c.post(
            "/api/projects/" + pid + "/sources",
            content=b"test",
            headers={**h, "X-File-Name": filename},
        ).status_code
        == 422
    )
    assert (
        c.post(
            "/api/projects/" + pid + "/sources",
            content=b"test",
            headers={"X-File-Name": "safe.md", "Origin": "http://testserver"},
        ).status_code
        == 403
    )


def test_office_pdf_locators_and_formula_preservation(tmp_path, monkeypatch):
    import subprocess

    original_run = subprocess.run

    def diagnosed_run(*args, **kwargs):
        result = original_run(*args, **kwargs)
        if result.returncode and "hub.extract" in args[0]:
            pytest.fail("Synthetic parser diagnostic: " + result.stderr.decode()[:5000])
        return result

    monkeypatch.setattr(subprocess, "run", diagnosed_run)
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from pypdf import PdfReader, PdfWriter

    doc = Document()
    doc.add_heading("Synthetic heading")
    doc.add_paragraph("Synthetic paragraph")
    doc.save(tmp_path / "test.docx")
    assert any(
        "Dokumentabschnitt" in c["locator"] and "paragraph" in c["content"]
        for c in isolated_parse(tmp_path / "test.docx", ".docx")["chunks"]
    )
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Synthetic slide"
    presentation.save(tmp_path / "test.pptx")
    assert isolated_parse(tmp_path / "test.pptx", ".pptx")["chunks"][0]["locator"] == "Folie 1"
    wb = Workbook()
    wb.active.append(["Name", "Amount"])
    wb.active.append(["A", 3])
    wb.active.append(["B", "=SUM(B2)"])
    wb.save(tmp_path / "test.xlsx")
    result = isolated_parse(tmp_path / "test.xlsx", ".xlsx")
    assert (
        result["chunks"][-1]["locator"] == "Sheet!A3:B3"
        and "=SUM(B2)" in result["chunks"][-1]["content"]
    )
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.encrypt("synthetic")
    writer.write(tmp_path / "locked.pdf")
    assert PdfReader(tmp_path / "locked.pdf").is_encrypted
    with pytest.raises(ValueError, match="Passwort"):
        parse(tmp_path / "locked.pdf", ".pdf")


def test_archive_active_content_and_bomb(tmp_path):
    for path in ["word/vbaProject.bin", "word/embeddings/file.exe", "../escape"]:
        archive = tmp_path / "unsafe.docx"
        with zipfile.ZipFile(archive, "w") as z:
            z.writestr(path, b"synthetic")
        with pytest.raises(ValueError):
            safe_archive(archive)
    with zipfile.ZipFile(tmp_path / "bomb.xlsx", "w", compression=zipfile.ZIP_DEFLATED) as z:
        z.writestr("sheet", "A" * 100000)
    with pytest.raises(ValueError):
        safe_archive(tmp_path / "bomb.xlsx")


def test_vault_bidirectional_rename_delete_restart_scope(knowledge):
    app, c, h, cfg, v = knowledge
    pid = project(c, h)
    mid = str(uuid.uuid4())
    body = {"id": mid, "title": "Synthetic", "content": "memory fact", "project_id": pid}
    first = c.post("/api/memories", json=body, headers=h)
    assert first.status_code == 200
    row = first.json()
    path = v / row["path"]
    assert path.is_file()
    assert c.post("/api/memories", json=body, headers=h).json()["id"] == mid
    changed = path.read_text().replace("memory fact", "edited in Obsidian")
    path.write_text(changed)
    assert c.post("/api/memories/sync", json={}, headers=h).status_code == 200
    rows = c.get("/api/memories?project_id=" + pid).json()
    assert rows[0]["revision"] == 3 and "edited in Obsidian" in rows[0]["content"]
    assert c.get("/api/memories").json() == []
    newpath = path.with_name("renamed.md")
    path.rename(newpath)
    c.post("/api/memories/sync", json={}, headers=h)
    assert c.get("/api/memories?project_id=" + pid).json()[0]["path"].endswith("renamed.md")
    newpath.unlink()
    c.post("/api/memories/sync", json={}, headers=h)
    assert c.get("/api/memories?project_id=" + pid).json() == []
    with app.state.sessions() as db:
        reconcile(db, v)
        db.commit()
        assert db.get(Memory, mid).deleted
        assert context(db, pid, "Obsidian") == context(db, pid, "no-match")
    assert c.get("/api/memories/" + mid + "/versions").json()


def test_memory_conflict_preserves_vault_and_readonly(knowledge):
    app, c, h, cfg, v = knowledge
    mid = str(uuid.uuid4())
    body = {"id": mid, "title": "Synthetic", "content": "base"}
    row = c.post("/api/memories", json=body, headers=h).json()
    path = v / row["path"]
    path.write_text(path.read_text().replace("base", "external revision"))
    body.update(content="app conflicting revision", revision=row["revision"])
    assert c.patch("/api/memories/" + mid, json=body, headers=h).status_code == 409
    assert "external revision" in path.read_text()
    assert c.get("/api/memories").json()[0]["conflict"]
    (v / "ordinary.md").write_text("# Human note\nprivate")
    c.post("/api/memories/sync", json={}, headers=h)
    human = next(m for m in c.get("/api/memories").json() if m["path"] == "ordinary.md")
    assert not human["writable"]
    assert (
        c.patch(
            "/api/memories/" + human["id"],
            json={**body, "id": human["id"], "revision": human["revision"]},
            headers=h,
        ).status_code
        == 409
    )


def test_vault_symlink_and_path_escape(tmp_path):
    root = tmp_path / "vault"
    root.mkdir()
    (root / "evil").symlink_to(tmp_path)
    for path in ["../escape", "/etc/passwd", ".obsidian/plugins/x", "evil/escape"]:
        with pytest.raises(ValueError):
            vault_path(root, path)


def test_web_private_query_blocked_before_network(monkeypatch):
    def fail(*args, **kwargs):
        pytest.fail("Private query reached DNS")

    monkeypatch.setattr("socket.getaddrinfo", fail)
    with pytest.raises(PermissionError):
        asyncio.run(
            search(Settings(), SearchInput(request_id=str(uuid.uuid4()), query="synthetic private"))
        )


def test_web_pins_and_rejects_endpoint(monkeypatch):
    monkeypatch.setattr(
        "socket.getaddrinfo", lambda *a, **kw: [(None, None, None, None, ("127.0.0.1", 8888))]
    )
    with pytest.raises(ValueError):
        target(Settings(search_url="http://search.local/search", search_allowed_ip="127.0.0.2"))
    assert (
        target(
            Settings(search_url="http://search.local:8888/search", search_allowed_ip="127.0.0.1")
        )[0]
        == "http://127.0.0.1:8888/search"
    )


def test_memory_explicit_conflict_resolution_bound_to_digest(knowledge):
    app, c, h, cfg, vault = knowledge
    identifier = str(uuid.uuid4())
    body = {"id": identifier, "title": "Synthetic", "content": "initial"}
    row = c.post("/api/memories", json=body, headers=h).json()
    path = vault / row["path"]
    path.write_text(path.read_text().replace("initial", "external"))
    body.update(content="chosen after comparison", revision=row["revision"])
    assert c.patch("/api/memories/" + identifier, json=body, headers=h).status_code == 409
    body["resolve_vault_sha256"] = digest("# Synthetic\n\nexternal")
    result = c.patch("/api/memories/" + identifier, json=body, headers=h)
    assert result.status_code == 200 and result.json()["conflict"] == ""
    assert "chosen after comparison" in path.read_text()
    assert len(c.get("/api/memories/" + identifier + "/versions").json()) >= 3


def test_memory_crash_intent_reconciles_without_overwrite(knowledge):
    app, c, h, cfg, vault = knowledge
    mid = str(uuid.uuid4())
    with app.state.sessions() as db:
        row = Memory(
            id=mid,
            path="AI/Memory/" + mid + ".md",
            writable=True,
            content="# Pending\n\napproved",
            base_content="",
            sha256=digest("# Pending\n\napproved"),
            revision=2,
            conflict="pending:" + digest("") + ":creating",
        )
        db.add(row)
        db.commit()
        reconcile(db, vault)
        db.commit()
        assert not row.deleted and row.conflict == ""
        assert (vault / row.path).is_file()
        (vault / row.path).write_text(
            (vault / row.path).read_text().replace("approved", "external")
        )
        row.content = "# Pending\n\nnext desired"
        row.sha256 = digest(row.content)
        row.conflict = "pending:" + digest("# Pending\n\napproved")
        db.commit()
        reconcile(db, vault)
        db.commit()
        assert row.conflict == "# Pending\n\nexternal"
        assert "external" in (vault / row.path).read_text()


def test_fetch_blocks_unapproved_unlisted_and_private_dns(monkeypatch):
    from hub.websearch import fetch_page

    def forbidden(*args, **kwargs):
        pytest.fail("Forbidden fetch reached DNS")

    monkeypatch.setattr("socket.getaddrinfo", forbidden)
    with pytest.raises(PermissionError):
        asyncio.run(fetch_page(Settings(), "https://example.org/", False))
    with pytest.raises(ValueError):
        asyncio.run(fetch_page(Settings(), "https://example.org/", True))
    monkeypatch.setattr(
        "socket.getaddrinfo", lambda *a, **kw: [(None, None, None, None, ("169.254.169.254", 443))]
    )
    cfg = Settings(web_domains_json='["example.org"]')
    with pytest.raises(ValueError):
        asyncio.run(fetch_page(cfg, "https://example.org/", True))


def test_explicit_remember_is_local_idempotent_and_not_from_sources(knowledge):
    app, c, h, cfg, vault = knowledge
    pid = project(c, h)
    chat = c.post("/api/conversations", json={"project_id": pid}, headers=h).json()["id"]
    r = c.post(
        "/api/conversations/" + chat + "/messages",
        json={"content": "Merke dir: Synthetische Prüffarbe ist grün."},
        headers=h,
    )
    assert "Im lokalen Vault gespeichert" in r.text and "succeeded" in r.text
    with app.state.sessions() as db:
        memories = list(db.scalars(select(Memory)))
        assert len(memories) == 1 and memories[0].project_id == pid
        assert memories[0].origin.startswith("chat:") and memories[0].label == "LOCAL_ONLY"
    c.post(
        "/api/conversations/" + chat + "/retry", json={"request_id": str(uuid.uuid4())}, headers=h
    )
    with app.state.sessions() as db:
        assert len(list(db.scalars(select(Memory)))) == 1
    assert app.state.providers.calls == 0


def test_extensionless_xlsx_original_and_reindex(knowledge, tmp_path):
    from openpyxl import Workbook

    app, c, h, cfg, v = knowledge
    wb = Workbook()
    wb.active.append(["A", 10.2])
    wb.active.append(["B", 2.3])
    path = tmp_path / "checksum_without_extension"
    wb.save(path)
    result = isolated_parse(path, ".xlsx")
    assert result["chunks"][0]["locator"] == "Sheet!A1:B1"
    pid = project(c, h)
    sid = upload(c, h, pid, "test.xlsx", path.read_bytes())
    assert wait(app, sid).status == "ready"
    checksum = wait(app, sid).sha256
    assert c.post(f"/api/projects/{pid}/sources/{sid}/retry", json={}, headers=h).status_code == 200
    row = wait(app, sid)
    assert row.status == "ready" and row.revision == 2 and row.sha256 == checksum


def test_global_vault_chat_requires_explicit_scope_and_stays_local(knowledge):
    app, c, h, cfg, v = knowledge
    (v / "global.md").write_text("# Global synthetic note\nGLOBAL_VAULT_FACT 42")
    c.post("/api/memories/sync", json={}, headers=h)
    pid = project(c, h)
    with app.state.sessions() as db:
        assert context(db, None, "GLOBAL_VAULT_FACT") == []
        assert not any(
            "GLOBAL_VAULT_FACT" in p.content for p in context(db, pid, "GLOBAL_VAULT_FACT")
        )
    chat = c.post("/api/conversations", json={"memory_scope": "global"}, headers=h).json()["id"]
    result = c.post(
        "/api/conversations/" + chat + "/messages", json={"content": "GLOBAL_VAULT_FACT"}, headers=h
    )
    assert "succeeded" in result.text
    detail = c.get("/api/conversations/" + chat).json()
    assert detail["memory_scope"] == "global" and detail["context_sources"]
    assert app.state.providers.stream_calls == [False]
    assert (
        c.post(
            "/api/conversations", json={"memory_scope": "global", "cloud_allowed": True}, headers=h
        ).status_code
        == 409
    )


def test_vault_disconnect_is_not_mistaken_for_deletions(knowledge, monkeypatch):
    app, c, h, cfg, vault = knowledge
    (vault / "stable.md").write_text("# Stable synthetic memory\nkeep on disconnect")
    c.post("/api/memories/sync", json={}, headers=h)
    before = c.get("/api/memories").json()

    def broken_walk(root, *, followlinks, onerror):
        onerror(OSError("synthetic disconnected filesystem"))
        return iter(())

    monkeypatch.setattr("hub.memory.os.walk", broken_walk)
    assert c.post("/api/memories/sync", json={}, headers=h).status_code == 409
    assert c.get("/api/memories").json() == before


def test_existing_memory_area_editable_only_with_folder_policy(knowledge):
    app, c, h, cfg, vault = knowledge
    (vault / "AI" / "Memory").mkdir(parents=True)
    (vault / "AI" / "Memory" / "human.md").write_text("# Human synthetic memory\nexisting fact")
    (vault / "ordinary.md").write_text("# Human general vault\nread only")
    assert c.post("/api/memories/sync", json={}, headers=h).status_code == 200
    rows = c.get("/api/memories").json()
    scoped = next(r for r in rows if r["path"].startswith("AI/Memory/"))
    assert scoped["writable"]
    assert not next(r for r in rows if r["path"] == "ordinary.md")["writable"]
    body = {
        "id": scoped["id"],
        "title": "Explicit revised memory",
        "content": "updated fact",
        "revision": scoped["revision"],
    }
    assert c.patch("/api/memories/" + scoped["id"], json=body, headers=h).status_code == 200
    assert "updated fact" in (vault / scoped["path"]).read_text()


def test_table_decimal_precision_and_ambiguous_formats_are_auditable(knowledge):
    app, c, h, cfg, v = knowledge
    pid = project(c, h)
    data = b'Value\n123456789012345678901234567890.12\n0.01\n"1,234.56"\nNaN\n'
    sid = upload(c, h, pid, "precision.csv", data)
    assert wait(app, sid).status == "ready"
    receipt = c.get(f"/api/projects/{pid}/sources/{sid}/sum?column=0").json()
    assert receipt["total"] == "123456789012345678901234567890.13"
    assert any(row["raw"] == "1,234.56" for row in receipt["skipped"])


def test_chat_web_requires_exact_user_query_grant_and_never_sends_private_context(
    knowledge, monkeypatch
):
    from dataclasses import replace

    app, c, h, cfg, vault = knowledge
    # Replace only the fixture worker configuration, never network/firewall policy.
    app.state.providers = FakeProviders(
        Route(capabilities=["web"], model_class="cloud", reason="public web fixture", answer=None)
    )
    import hub.api as api_module

    original_execute = api_module.execute_run

    async def configured_execute(factory, settings, providers, run_id):
        await original_execute(
            factory,
            replace(settings, search_url="http://synthetic.local/search"),
            providers,
            run_id,
        )

    monkeypatch.setattr(api_module, "execute_run", configured_execute)
    sent = []

    async def web_fixture(settings, body):
        assert body.public_query_approved
        sent.append(body.query)
        return [
            {
                "url": "https://example.org/synthetic",
                "title": "Synthetic public evidence",
                "content": "PUBLIC_WEB_FACT",
                "retrieved_at": 1.0,
                "provider": "searxng",
            }
        ]

    monkeypatch.setattr("hub.websearch.search", web_fixture)
    cid = c.post("/api/conversations", json={}, headers=h).json()["id"]
    first = c.post(
        "/api/conversations/" + cid + "/messages",
        json={"content": "PRIVATE_HISTORY_MARKER"},
        headers=h,
    )
    assert "failed" in first.text and sent == []
    second = c.post(
        "/api/conversations/" + cid + "/messages",
        json={"content": "Exact public query", "web_search_allowed": True},
        headers=h,
    )
    assert "succeeded" in second.text and sent == ["Exact public query"]
    assert "PRIVATE_HISTORY_MARKER" not in sent[0]
    assert app.state.providers.stream_calls == [False]
    detail = c.get("/api/conversations/" + cid).json()
    assert any(
        s["reference"].startswith("WEBQUELLE https://example.org/")
        for s in detail["context_sources"]
    )
    retried = c.post(
        "/api/conversations/" + cid + "/retry", json={"request_id": str(uuid.uuid4())}, headers=h
    )
    assert "failed" in retried.text and len(sent) == 1


def test_vault_file_transport_failure_rolls_back_entire_scan(knowledge, monkeypatch):
    import errno

    app, c, h, cfg, vault = knowledge
    note = vault / "stable.md"
    note.write_text("# Stable synthetic memory\noriginal")
    assert c.post("/api/memories/sync", json={}, headers=h).status_code == 200
    before = c.get("/api/memories").json()
    import os

    original_open = os.open

    def failed_open(path, *args, **kwargs):
        if path == note.name:
            raise OSError(errno.EIO, "synthetic lost SFTP connection")
        return original_open(path, *args, **kwargs)

    monkeypatch.setattr(os, "open", failed_open)
    assert c.post("/api/memories/sync", json={}, headers=h).status_code == 409
    assert c.get("/api/memories").json() == before


def test_global_vault_matches_before_recent_limit_and_avoids_unrelated_context(knowledge):
    app, c, h, cfg, vault = knowledge
    with app.state.sessions() as db:
        db.add(
            Memory(
                path="older.md",
                content="UNIQUE_OLDER_VAULT_FACT 42",
                sha256=hashlib.sha256(b"UNIQUE_OLDER_VAULT_FACT 42").hexdigest(),
                updated=1,
            )
        )
        db.add_all(
            Memory(
                path=f"newer-{i}.md",
                content="Unrelated synthetic note",
                sha256=hashlib.sha256(b"Unrelated synthetic note").hexdigest(),
                updated=10 + i,
            )
            for i in range(501)
        )
        db.commit()
        parts = context(db, None, "UNIQUE_OLDER_VAULT_FACT", global_memory=True)
        assert len(parts) == 1 and "older.md" in parts[0].content
        with pytest.raises(ValueError, match="Keine passenden"):
            context(db, None, "NO_MATCH_MARKER", global_memory=True)


def test_existing_empty_directory_does_not_satisfy_required_vault_mount(tmp_path):
    from hub.memory import validate_mount

    (tmp_path / "AI/Memory").mkdir(parents=True)
    with pytest.raises(ValueError, match="Mount nicht verfügbar"):
        validate_mount(tmp_path, "fuse.sshfs")


def test_vault_descriptor_read_denies_symlink_and_oversize(tmp_path):
    from hub.memory import MAX_NOTE, read_note

    root = tmp_path / "vault"
    root.mkdir()
    (root / "folder").mkdir()
    (root / "folder" / "safe.md").write_text("# Synthetic safe note")
    assert read_note(root, "folder/safe.md") == "# Synthetic safe note"
    outside = tmp_path / "private.md"
    outside.write_text("Outside must never be read")
    (root / "link.md").symlink_to(outside)
    (root / "linked-folder").symlink_to(tmp_path)
    for path in ["link.md", "linked-folder/private.md"]:
        with pytest.raises(OSError):
            read_note(root, path)
    for path in ["../private.md", str(outside), ".obsidian/config.md"]:
        with pytest.raises(ValueError):
            read_note(root, path)
    (root / "large.md").write_bytes(b"x" * (MAX_NOTE + 1))
    with pytest.raises(ValueError):
        read_note(root, "large.md")


def test_duplicate_content_does_not_probe_every_existing_candidate_path(knowledge, monkeypatch):
    import hub.memory as memory

    app, c, h, cfg, vault = knowledge
    for i in range(30):
        (vault / f"duplicate-{i}.md").write_text("# Identical synthetic template")
    original = memory.vault_path
    probed = []

    def tracked(root, relative):
        probed.append(relative)
        return original(root, relative)

    monkeypatch.setattr(memory, "vault_path", tracked)
    assert c.post("/api/memories/sync", json={}, headers=h).status_code == 200
    assert len(c.get("/api/memories").json()) == 30
    assert not any(p.startswith("duplicate-") for p in probed)
    (vault / "duplicate-0.md").rename(vault / "renamed-template.md")
    before = c.get("/api/memories").json()
    expected = next(m["id"] for m in before if m["path"] == "duplicate-0.md")
    assert c.post("/api/memories/sync", json={}, headers=h).status_code == 200
    after = c.get("/api/memories").json()
    assert next(m["id"] for m in after if m["path"] == "renamed-template.md") == expected


def test_vault_scan_opens_shared_parent_once_and_rejects_fifo(knowledge, monkeypatch):
    import os

    from hub.memory import read_note

    app, c, h, cfg, vault = knowledge
    for i in range(20):
        (vault / f"shared-{i}.md").write_text(f"# Synthetic {i}")
    original = os.open
    roots = []

    def opened(path, *args, **kwargs):
        if path == vault:
            roots.append(path)
        return original(path, *args, **kwargs)

    monkeypatch.setattr(os, "open", opened)
    assert c.post("/api/memories/sync", json={}, headers=h).status_code == 200
    assert len(c.get("/api/memories").json()) == 20
    assert len(roots) == 1
    os.mkfifo(vault / "pipe.md")
    with pytest.raises(ValueError, match="reguläre Datei"):
        read_note(vault, "pipe.md")
