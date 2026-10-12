import shutil
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from hub.artifact_render import _spreadsheet
from hub.artifact_specs import SpreadsheetPlan
from hub.spreadsheet_features import cache_formula_results
from openpyxl import load_workbook
from pydantic import ValidationError


def plan(rows, **kwargs):
    return SpreadsheetPlan.model_validate(
        {
            "kind": "xlsx",
            "title": "Calculation",
            "sheets": [
                {"name": "Input", "columns": ["Item", "Amount", "Result"], "rows": rows, **kwargs}
            ],
        }
    )


@pytest.mark.parametrize(
    "formula",
    [
        '=WEBSERVICE("https://invalid.example")',
        '=HYPERLINK("file:///secret")',
        "='[private.xlsx]Sheet1'!A2",
        "=Input!D2",
        "=A999",
        "=SUM(A:A)",
        '=INDIRECT("B2")',
        "=A2 B2",
        "=SUM({1,2})",
        "=SUM(A2",
    ],
)
def test_formula_boundaries(formula):
    with pytest.raises(ValidationError):
        plan([["x", 2, {"formula": formula}]])


def test_formula_graph_cross_sheet_and_cycles():
    value = SpreadsheetPlan.model_validate(
        {
            "kind": "xlsx",
            "title": "Budget",
            "sheets": [
                {"name": "Cost base", "columns": ["ID", "Amount"], "rows": [["DC-001", -120]]},
                {
                    "name": "Report",
                    "columns": ["Item", "Value"],
                    "rows": [
                        ["Cost", {"formula": "='Cost base'!$B$2"}],
                        ["Check", {"formula": '=IF(ROUND(B2+120,2)=0,"Ties","Does not tie")'}],
                    ],
                },
            ],
        }
    )
    assert value.sheets[1].rows[0][1].formula.startswith("=")
    for rows in [
        [["x", {"formula": "=C2"}, {"formula": "=B2"}]],
        [["x", 2, {"formula": "=C2"}]],
    ]:
        with pytest.raises(ValidationError, match="Zirkelbezug"):
            plan(rows)


def test_native_charts_formats_totals_and_literals(tmp_path, monkeypatch):
    monkeypatch.setattr("hub.artifact_render.cache_formula_results", lambda path: None)
    value = plan(
        [["=literal", -120, {"formula": "=ABS(B2)"}], ["other", -80, {"formula": "=ABS(B3)"}]],
        charts=[
            {
                "type": "bar",
                "title": "Costs",
                "categories": "A2:A3",
                "series": ["B1:B3"],
                "anchor": "E2",
            },
            {
                "type": "scatter",
                "title": "Comparison",
                "categories": "B2:B3",
                "series": ["C1:C3"],
                "anchor": "E22",
            },
        ],
        number_formats={"B": "currency", "C": "decimal"},
        totals={"B": "sum"},
        highlight_cells=["B3"],
        column_groups=[{"start": "C", "end": "C"}],
        no_stripes=True,
    )
    path = tmp_path / "artifact.xlsx"
    _spreadsheet(value, path)
    book = load_workbook(path)
    sheet = book["Input"]
    assert sheet["A2"].data_type == "s"
    assert sheet["C2"].data_type == "f"
    assert sheet["B4"].value == "=SUBTOTAL(109,B2:B3)"
    assert len(sheet._charts) == 2
    table = sheet.tables["tbl_data_1"]
    assert table.tableColumns[1].totalsRowFunction == "sum"
    assert table.autoFilter.ref == "A1:C3"
    assert not table.tableStyleInfo.showRowStripes
    assert sheet["B3"].fill.fgColor.rgb.endswith("FFF2CC")
    assert sheet.column_dimensions["C"].hidden
    assert sheet["B2"].font.color.rgb.endswith("0000FF")
    assert sheet["C2"].font.color.rgb.endswith("000000")
    assert sheet["B2"].number_format.startswith('_("€"')


@pytest.mark.parametrize(
    "extensions",
    [
        {"charts": [{"type": "bar", "title": "Chart", "categories": "A2:A3", "series": ["B1:B2"]}]},
        {
            "charts": [
                {
                    "type": "pie",
                    "title": "Chart",
                    "categories": "A2:A2",
                    "series": ["B1:B2", "C1:C2"],
                }
            ]
        },
        {"highlight_cells": ["A3"]},
        {"column_groups": [{"start": "C", "end": "B"}]},
        {"totals": {"D": "sum"}},
    ],
)
def test_invalid_extensions(extensions):
    with pytest.raises(ValidationError):
        plan([["x", 2, 3]], **extensions)


def test_cache_injection_preserves_native_zip(tmp_path, monkeypatch):
    monkeypatch.setattr("hub.artifact_render.cache_formula_results", lambda path: None)
    path = tmp_path / "artifact.xlsx"
    _spreadsheet(
        plan(
            [["x", 2, {"formula": "=B2*3"}]],
            charts=[
                {"type": "bar", "title": "Native chart", "categories": "A2:A2", "series": ["B1:B2"]}
            ],
        ),
        path,
    )
    with zipfile.ZipFile(path) as source:
        before = {name: source.read(name) for name in source.namelist()}

    def calculate(command, **kwargs):
        target = Path(command[command.index("--outdir") + 1]) / path.name
        ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        with zipfile.ZipFile(path) as source, zipfile.ZipFile(target, "w") as output:
            for name in source.namelist():
                data = source.read(name)
                if name == "xl/worksheets/sheet1.xml":
                    root = ET.fromstring(data)
                    for cell in root.iter(ns + "c"):
                        if cell.attrib.get("r") == "C2":
                            cell.find(ns + "v").text = "6"
                    data = ET.tostring(root)
                output.writestr(name, data)
        import subprocess

        return subprocess.CompletedProcess(command, 0)

    monkeypatch.setattr("subprocess.run", calculate)
    cache_formula_results(path)
    assert load_workbook(path, data_only=True)["Input"]["C2"].value == 6
    assert load_workbook(path, data_only=False)["Input"]["C2"].value == "=B2*3"
    with zipfile.ZipFile(path) as source:
        for name, data in before.items():
            if name != "xl/worksheets/sheet1.xml":
                assert source.read(name) == data


@pytest.mark.skipif(not shutil.which("soffice"), reason="Offline Calc not installed locally")
def test_actual_offline_calc_formula_results(tmp_path):
    path = tmp_path / "artifact.xlsx"
    _spreadsheet(
        plan(
            [
                ["Revenue", 240, {"formula": "=B2*1.2"}],
                ["Costs", -120, {"formula": "=SUM(B2:B3)"}],
                ["Check", None, {"formula": '=IF(C3=120,"Ties","Mismatch")'}],
            ],
            totals={"B": "sum"},
            number_formats={"B": "currency"},
            charts=[
                {
                    "type": "bar",
                    "title": "Revenue and costs",
                    "categories": "A2:A3",
                    "series": ["B1:B3"],
                }
            ],
        ),
        path,
    )
    book = load_workbook(path, data_only=True)
    assert book["Input"]["C2"].value == 288
    assert book["Input"]["C3"].value == 120
    assert book["Input"]["C4"].value == "Ties"
    assert book["Input"]["B5"].value == 120
    native = load_workbook(path, data_only=False)
    assert len(native["Input"]._charts) == 1
    assert native["Input"].tables["tbl_data_1"].tableStyleInfo.name == "TableStyleLight8"


@pytest.mark.skipif(not shutil.which("soffice"), reason="Offline Calc not installed locally")
def test_actual_cross_sheet_financial_and_date_formulas(tmp_path):
    value = SpreadsheetPlan.model_validate(
        {
            "kind": "xlsx",
            "title": "Financial model",
            "sheets": [
                {
                    "name": "Cost base",
                    "columns": ["ID", "Value", "Category"],
                    "rows": [["DC-001", 120, "Confirmed"], ["DC-002", 80, "Open"]],
                },
                {
                    "name": "Report",
                    "columns": ["Check", "Result"],
                    "number_formats": {"B": "decimal"},
                    "rows": [
                        [
                            "Conditional total",
                            {
                                "formula": "=SUMIF('Cost base'!C2:C3,\"Confirmed\","
                                "'Cost base'!B2:B3)"
                            },
                        ],
                        [
                            "Lookup",
                            {
                                "formula": "=INDEX('Cost base'!B2:B3,"
                                "MATCH(\"DC-002\",'Cost base'!A2:A3,0))"
                            },
                        ],
                        [
                            "Weights",
                            {"formula": "=SUMPRODUCT('Cost base'!B2:B3,'Cost base'!B2:B3)"},
                        ],
                        ["Payment", {"formula": "=ROUND(PMT(0,12,1200),2)"}],
                        [
                            "Control",
                            {"formula": '=IF(AND(B2=120,B3=80,B4=20800),"Ties","Mismatch")'},
                        ],
                    ],
                },
                {
                    "name": "Dates",
                    "columns": ["Item", "Date"],
                    "number_formats": {"B": "date"},
                    "rows": [
                        ["Start", {"formula": "=DATE(2026,1,1)"}],
                        ["Blank", {"formula": '=IF(1=1,"",0)'}],
                    ],
                },
            ],
        }
    )
    path = tmp_path / "artifact.xlsx"
    _spreadsheet(value, path)
    cached = load_workbook(path, data_only=True)
    assert [cached["Report"][f"B{number}"].value for number in range(2, 7)] == [
        120,
        80,
        20800,
        -100,
        "Ties",
    ]
    assert cached["Dates"]["B2"].value.year == 2026
    native = load_workbook(path, data_only=False)
    assert native["Report"]["B2"].font.color.rgb.endswith("008000")
    assert native["Report"]["B2"].value.startswith("=SUMIF")


def test_total_cycle_and_duplicate_headers_rejected():
    with pytest.raises(ValidationError, match="Zirkelbezug"):
        plan([["x", {"formula": "=B3"}, 1]], totals={"B": "sum"})
    with pytest.raises(ValidationError, match="eindeutig"):
        SpreadsheetPlan.model_validate(
            {
                "kind": "xlsx",
                "title": "Bad headers",
                "sheets": [
                    {
                        "name": "Input",
                        "columns": ["Amount", "amount"],
                        "rows": [[1, {"formula": "=A2"}]],
                    }
                ],
            }
        )


def test_legacy_duplicate_headers_retain_canonical_shape():
    value = SpreadsheetPlan.model_validate(
        {
            "kind": "xlsx",
            "title": "Legacy",
            "sheets": [{"name": "Input", "columns": ["Amount", "amount"], "rows": [[1, 2]]}],
        }
    )
    assert value.model_dump(mode="json")["sheets"][0] == {
        "name": "Input",
        "columns": ["Amount", "amount"],
        "rows": [[1, 2]],
    }


def test_dependency_budget_rejected_during_graph_build():
    with pytest.raises(ValidationError, match="Arbeitsmappenlimit"):
        SpreadsheetPlan.model_validate(
            {
                "kind": "xlsx",
                "title": "Bounded",
                "sheets": [
                    {
                        "name": "Source",
                        "columns": [f"Column {number}" for number in range(10)],
                        "rows": [[0] * 10 for _ in range(500)],
                    },
                    {
                        "name": "Report",
                        "columns": ["Value"],
                        "rows": [[{"formula": "=SUM(Source!A1:J501)"}] for _ in range(45)],
                    },
                ],
            }
        )


def test_future_function_prefix_changes_only_function_tokens():
    from hub.spreadsheet_features import native_formula

    assert native_formula('=CONCAT("XLOOKUP(",IFS(B2>0,"yes",TRUE,"no"))') == (
        '=_xlfn.CONCAT("XLOOKUP(",_xlfn.IFS(B2>0,"yes",TRUE,"no"))'
    )
    assert native_formula("=SUM(B2:B3)") == "=SUM(B2:B3)"


@pytest.mark.skipif(not shutil.which("soffice"), reason="Offline Calc not installed locally")
def test_future_functions_native_ooxml_and_real_calc_caches(tmp_path):
    value = plan(
        [
            ["one", 12, {"formula": '=XLOOKUP("one",A2:A4,B2:B4)'}],
            ["two", 24, {"formula": '=IFS(B3=24,"Matches",TRUE,"Missing")'}],
            ["three", 36, {"formula": '=CONCAT(A2," / ",A3)'}],
        ]
    )
    path = tmp_path / "artifact.xlsx"
    _spreadsheet(value, path)
    cached = load_workbook(path, data_only=True)
    assert [cached["Input"][f"C{n}"].value for n in range(2, 5)] == [12, "Matches", "one / two"]
    native = load_workbook(path, data_only=False)
    assert native["Input"]["C2"].value == '=_xlfn.XLOOKUP("one",A2:A4,B2:B4)'
    assert native["Input"]["C3"].value.startswith("=_xlfn.IFS(")
    assert native["Input"]["C4"].value.startswith("=_xlfn.CONCAT(")
    assert value.sheets[0].rows[0][2].formula.startswith("=XLOOKUP(")
    with zipfile.ZipFile(path) as source:
        xml = source.read("xl/worksheets/sheet1.xml")
        assert b"_xlfn.XLOOKUP(" in xml and b"_xlfn.IFS(" in xml and b"_xlfn.CONCAT(" in xml


@pytest.mark.parametrize(
    "formula",
    [
        '=XLOOKUP("one",A2:A3,B2:B2)',
        '=XLOOKUP("one",A2:B3,B2:C3)',
        '=XLOOKUP("one",A2:A3,B2:B3,"missing",-1)',
        '=XLOOKUP("one",A2:A3,B2:B3,"missing",0,-1)',
        '=XLOOKUP("one",A2:A3,B2:C2)',
    ],
)
def test_xlookup_unsupported_modes_and_array_shapes_rejected(formula):
    with pytest.raises(ValidationError):
        plan([["one", 12, {"formula": formula}], ["two", 24, None]])


def test_xlookup_calculation_copy_lowering_preserves_native_formula():
    from hub.spreadsheet_features import calculation_formula

    native = '=_xlfn.XLOOKUP("*?~",A2:A4,B2:B4,"Missing")'
    lowered = calculation_formula(native)
    assert "_xlfn.XLOOKUP" not in lowered
    assert lowered.startswith("=IF(ISNA(MATCH(")
    assert "SUBSTITUTE" in lowered and '"~*"' in lowered and '"~?"' in lowered
    assert "INDEX(B2:B4," in lowered and lowered.endswith(",1))")
    horizontal = calculation_formula('=XLOOKUP("one",A2:C2,A3:C3)')
    assert "INDEX(A3:C3,1,MATCH(" in horizontal


@pytest.mark.skipif(not shutil.which("soffice"), reason="Offline Calc not installed locally")
def test_lowered_xlookup_literal_wildcards_missing_and_horizontal(tmp_path):
    value = SpreadsheetPlan.model_validate(
        {
            "kind": "xlsx",
            "title": "Exact lookup",
            "sheets": [
                {
                    "name": "Source",
                    "columns": ["Key", "Amount"],
                    "rows": [["*one", 12], ["one", 24]],
                },
                {
                    "name": "Horizontal",
                    "columns": ["First", "Second", "Third"],
                    "rows": [["first", "second", "third"], [12, 24, 36]],
                },
                {
                    "name": "Report",
                    "columns": ["Check", "Result"],
                    "rows": [
                        [
                            "Literal wildcard",
                            {"formula": '=XLOOKUP("*one",Source!A2:A3,Source!B2:B3)'},
                        ],
                        [
                            "Missing",
                            {"formula": '=XLOOKUP("absent",Source!A2:A3,Source!B2:B3,"Missing")'},
                        ],
                        [
                            "Horizontal",
                            {"formula": '=XLOOKUP("second",Horizontal!A2:C2,Horizontal!A3:C3)'},
                        ],
                    ],
                },
            ],
        }
    )
    path = tmp_path / "artifact.xlsx"
    _spreadsheet(value, path)
    cached = load_workbook(path, data_only=True)
    assert [cached["Report"][f"B{n}"].value for n in range(2, 5)] == [12, "Missing", 24]
    assert load_workbook(path, data_only=False)["Report"]["B2"].value.startswith("=_xlfn.XLOOKUP(")


@pytest.mark.parametrize("kind", ["bar", "line", "area", "pie", "scatter"])
def test_chart_data_only_series_keeps_first_value_and_header_title(tmp_path, kind):
    value = plan(
        [["First", 10, 1], ["Second", 20, 2], ["Third", 30, 3]],
        charts=[
            {
                "type": kind,
                "title": "All values",
                "categories": "A2:A4" if kind != "scatter" else "C2:C4",
                "series": ["B2:B4"],
            }
        ],
    )
    path = tmp_path / "artifact.xlsx"
    _spreadsheet(value, path)
    series = load_workbook(path)["Input"]._charts[0].series[0]
    values = series.yVal if kind == "scatter" else series.val
    assert values.numRef.f == "'Input'!$B$2:$B$4"
    assert series.tx.strRef.f == "'Input'!$B$1"


@pytest.mark.parametrize("data_range", ["B1:B3", "B2:B3", "B2:B5"])
def test_chart_mismatched_ranges_still_rejected(data_range):
    with pytest.raises(ValidationError):
        plan(
            [["First", 10, 1], ["Second", 20, 2], ["Third", 30, 3]],
            charts=[
                {
                    "type": "bar",
                    "title": "Invalid",
                    "categories": "A2:A4",
                    "series": [data_range],
                }
            ],
        )
